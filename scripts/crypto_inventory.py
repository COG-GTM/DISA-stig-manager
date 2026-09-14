#!/usr/bin/env python3
"""Cryptographic inventory scanner for this repository.

Walks the working tree, detects cryptographic primitives, protocol
configuration, key material, crypto libraries, runtime pins and secret
locations, and writes:

  <out>/crypto-inventory.json
  <out>/Cryptographic-Inventory.xlsx          (requires openpyxl)
  <out>/PQC-Readiness-Assessment.md           (rendered from the template in <out>)

The scanner is heuristic. It matches source text with regular expressions and
inspects a small window of surrounding lines. It does not execute application
code. Limits are listed in the `Method` sheet and in METHOD_LIMITS below.

Secret handling: for findings in the "Secret" category only the path, line and
type are recorded. Matched text is never written to any artifact.

Standard library only, plus optional:
  openpyxl      -> XLSX output
  cryptography  -> certificate / public key parsing (falls back to `openssl`)
  node          -> local runtime probe (versions, PQC algorithm availability)
"""

from __future__ import annotations

import argparse
import base64
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
from collections import Counter, OrderedDict
from pathlib import Path

# --------------------------------------------------------------------------- #
# Constants
# --------------------------------------------------------------------------- #

Q_SHOR = "Quantum-vulnerable (Shor)"
Q_GROVER = "Symmetric — Grover-affected, adequate at ≥256-bit"
Q_WEAK = "Deprecated/weak regardless of PQC"
Q_PROTO = "Protocol/configuration — inherits class of negotiated algorithms"
Q_NA = "Not applicable (no quantum-relevant primitive)"

CLASS_ORDER = [Q_SHOR, Q_WEAK, Q_GROVER, Q_PROTO, Q_NA]

CNSA_KEM = "ML-KEM-1024 (FIPS 203) for key establishment; hybrid X25519MLKEM768 as interim"
CNSA_SIG = "ML-DSA-87 (FIPS 204) or SLH-DSA (FIPS 205) for signatures"
CNSA_SW_SIG = "ML-DSA (FIPS 204) or LMS/XMSS (SP 800-208) for software signing"
CNSA_TLS = "TLS 1.3 with ML-KEM key establishment and ML-DSA certificates when available"
CNSA_HASH = "SHA-384 or SHA-512 (already SHA-2; SHA-256 acceptable per CNSA 2.0 where interoperability requires)"
CNSA_SYM = "AES-256"
CNSA_NONE = "None required"
CNSA_REMOVE = "Retire; replace with SHA-384/SHA-512 or an approved algorithm"

AGILITY_CONFIG = "Configurable"
AGILITY_LIB = "Library-bound"
AGILITY_HARD = "Hardcoded"
AGILITY_NA = "N/A"

EXCLUDE_DIR_NAMES = {
    "node_modules", "dist", ".git", "_build", ".nyc_output", "coverage",
    "mochawesome-report", "__pycache__", ".pytest_cache", ".ruff_cache", ".mypy_cache",
    ".venv", "venv", "uploads",
}
# Vendored third-party code. Relative posix paths.
EXCLUDE_DIR_PATHS = {"client/src/ext", "docs/security/crypto-inventory", "test/api/form-data-files"}
# This scanner and its outputs (they contain the patterns being searched for).
EXCLUDE_FILE_PATHS = {"scripts/crypto_inventory.py"}

TEXT_EXTENSIONS = {
    ".js", ".mjs", ".cjs", ".ts", ".json", ".yml", ".yaml", ".sh", ".bat",
    ".html", ".py", ".rst", ".md", ".txt", ".csv", ".env", ".example",
    ".conf", ".cnf", ".toml", ".ini", ".properties", ".sql",
}
TEXT_BASENAMES = {"Dockerfile", ".gitignore", ".dockerignore", ".npmrc", ".env"}
SKIP_FILE_RES = [re.compile(p) for p in (
    r"\.min\.js$", r"jsonview\.bundle\.js$", r"package-lock\.json$",
    r"\.lock$", r"stigman-asd-full\.csv$", r"\.ckl$",
)]
LOCKFILE_NAMES = {"package-lock.json", "npm-shrinkwrap.json"}
CERT_EXTENSIONS = {".pem", ".crt", ".cer", ".key", ".der", ".p12", ".pfx", ".jks", ".csr"}
MAX_TEXT_BYTES = 2 * 1024 * 1024
CALL_WINDOW = 4  # lines a multi-line call may span, including its first line
PEM_MAX_LINES = 400  # upper bound on the lines one PEM block may span

WINDOW = 12  # lines of context inspected before/after a match

# Libraries that carry a cryptographic role. Version comes from lockfiles.
CRYPTO_LIBRARIES = OrderedDict([
    ("jsonwebtoken", {
        "role": "JWT/JWS parse and signature verification (API access tokens)",
        "pqc": "None. Supports HS*, RS*, PS*, ES* only; no EdDSA, no ML-DSA. Accepted algorithms are chosen from the key type when `algorithms` is not passed.",
        "notes": "Used in api/source/utils/auth.js and test/utils/mockOidc.js.",
    }),
    ("jwks-rsa", {
        "role": "JWKS client (declared dependency)",
        "pqc": "None. RSA/EC JWKS import only.",
        "notes": "Declared in api/source/package.json; no `require('jwks-rsa')` found in application code. The application uses its own JWKSCache (api/source/utils/jwksCache.js).",
    }),
    ("jose", {
        "role": "JOSE/JWK library (transitive via jwks-rsa)",
        "pqc": "None in the resolved 4.x line.",
        "notes": "Not required directly by application code.",
    }),
    ("mysql2", {
        "role": "MySQL client; TLS to the database through Node tls",
        "pqc": "Inherits Node/OpenSSL TLS capabilities. Hybrid ML-KEM key exchange requires Node with OpenSSL 3.5+ and a MySQL server built against OpenSSL 3.5+.",
        "notes": "TLS enabled only when STIGMAN_DB_TLS_* variables are set (api/source/service/utils.js).",
    }),
    ("undici", {
        "role": "HTTP client for OIDC discovery; TLS to the identity provider",
        "pqc": "Inherits Node/OpenSSL TLS capabilities.",
        "notes": "Custom CA supplied through STIGMAN_OIDC_CA_CERTS (api/source/utils/auth.js).",
    }),
    ("ws", {
        "role": "WebSocket server (log stream); relies on the HTTP(S) server TLS",
        "pqc": "Inherits Node/OpenSSL TLS capabilities.",
        "notes": "No independent cryptography.",
    }),
    ("express", {
        "role": "HTTP framework; no cryptography of its own",
        "pqc": "N/A",
        "notes": "TLS is provided by node:https in api/source/bootstrap/server.js.",
    }),
    ("jszip", {
        "role": "ZIP read/write; CRC-32 only (non-cryptographic)",
        "pqc": "N/A",
        "notes": "Not a cryptographic dependency.",
    }),
    ("archiver", {
        "role": "Archive creation; CRC-32 only (non-cryptographic)",
        "pqc": "N/A",
        "notes": "Not a cryptographic dependency.",
    }),
    ("node-forge", {"role": "TLS/PKI in JavaScript", "pqc": "None", "notes": ""}),
    ("bcrypt", {"role": "Password hashing", "pqc": "N/A (symmetric/KDF)", "notes": ""}),
    ("bcryptjs", {"role": "Password hashing", "pqc": "N/A (symmetric/KDF)", "notes": ""}),
    ("argon2", {"role": "Password hashing", "pqc": "N/A (symmetric/KDF)", "notes": ""}),
    ("crypto-js", {"role": "JavaScript crypto primitives", "pqc": "None", "notes": ""}),
    ("openpgp", {"role": "OpenPGP", "pqc": "Draft PQC support only", "notes": ""}),
    ("keycloak-js", {"role": "OIDC client adapter", "pqc": "Depends on IdP", "notes": ""}),
    ("oidc-client-ts", {"role": "OIDC client", "pqc": "Depends on IdP", "notes": ""}),
    ("helmet", {"role": "HTTP security headers (HSTS)", "pqc": "N/A", "notes": ""}),
])

METHOD_CHECKS = [
    "Regex scan of text files (extensions: " + ", ".join(sorted(TEXT_EXTENSIONS)) + "; plus Dockerfile, .gitignore) for asymmetric, symmetric, hash, KDF, TLS, JWT/JWS/JWKS, randomness and secret patterns.",
    "Node.js crypto API calls are matched by name: generateKeyPair/Sync (rsa, rsa-pss, ec, ed25519, ed448, x25519, x448, dsa, dh), createECDH, diffieHellman, computeSecret, createDiffieHellman/Group, getDiffieHellman, createSign/createVerify, crypto.sign/verify, createCipheriv/createDecipheriv, and Web Crypto subtle.* calls with a public-key algorithm name. Curve, prime length and cipher name are taken from the literal arguments of the same call when present. Call rules are matched against the call's first line plus the next %d lines, so an argument list that continues on the following line is still inventoried once, at the line where the call starts." % (CALL_WINDOW - 1),
    "Context window of ±%d lines is inspected to infer key sizes (modulusLength, JWK `n` length), purpose (PKCE, kid derivation, attachment metadata), and presence/absence of TLS or JWT verification options." % WINDOW,
    "Certificate and key files by extension (%s) are parsed with `cryptography` or the `openssl` CLI. PEM certificates give subject/issuer attribute types, key algorithm/size, signature algorithm and validity; PEM public keys and unencrypted private keys give key algorithm and size only (public parameters; private components are never read into the output). Files without a PEM block are tried as DER certificate, then DER private key, then DER public key, and get a KEYMAT-DER-* inventory row; keystores and files that parse as none of these get a KEYMAT-FILE-UNPARSED row (priority 2) and are recorded as not parsed. Encrypted keys and unreadable files are recorded as not parsed." % ", ".join(sorted(CERT_EXTENSIONS)),
    "PEM blocks (certificate, public key, private key) found in any scanned text file are parsed the same way and the parsed algorithm, size, signature algorithm and validity are copied into the KEYMAT-PEM-* inventory row; RSA/DSA below 2048 bits and SHA-1/MD5 certificate signatures are classified as deprecated/weak. Blocks that cannot be parsed keep the block-type classification and say so in the Mode column. Public keys are also parsed from JWKS `x5c` arrays and TUF/Notary root metadata (root.json).",
    "package.json and package-lock.json files are read for declared and resolved versions of libraries with a cryptographic role.",
    "Dockerfiles, GitHub Actions workflows and pkg build configuration are read for Node.js runtime pins.",
    "If `node` is on PATH, the local runtime is probed for Node/OpenSSL versions, default TLS versions and ML-KEM/ML-DSA availability. This describes the scanning host, not the deployed container.",
    "Secrets: only path, line and type are recorded; matched values are never written.",
    "Certificate subject/issuer: only attribute types and self-signed status are recorded unless --dn-values is given.",
    "Absence checks: if no HSTS header and no explicit TLS option (minVersion, maxVersion, ciphers, rejectUnauthorized, secureOptions, honorCipherOrder, ecdhCurve) is found anywhere in scanned files, one row per absence is added and anchored to the TLS server setup line.",
]

METHOD_LIMITS = [
    "Static text matching only. Dynamic algorithm selection (for example the algorithm list jsonwebtoken derives from the key type) is inferred from library source knowledge, not observed at runtime.",
    "Key sizes are recorded only where they appear in source (modulusLength), can be derived from embedded key material (JWK `n`, SPKI), or are stated in comments. Keys supplied at deployment time (TLS certificates, IdP signing keys) are not visible to the scanner.",
    "TLS protocol versions and cipher suites negotiated at runtime depend on the Node.js/OpenSSL build of the container image and on peers (reverse proxy, MySQL server, IdP, browsers). The scanner records what the repository configures and flags what it does not configure.",
    "Vendored third-party code (client/src/ext), minified bundles, lockfiles (except for version extraction), XCCDF/CKL STIG content fixtures (test/api/form-data-files, *.xml, *.ckl), generated docs, this scanner and its own output (the canonical docs/security/crypto-inventory directory, any in-repository --out directory, and the concrete JSON/XLSX/report/template output paths, so that `--out .` does not inventory the previous run) are not pattern-scanned.",
    "Text PEM files with certificate/key extensions are parsed for metadata and also run through the pattern rules, so a committed private key appears both in the certificate table and as a Secret row. A PEM block may span at most %d lines; longer blocks are recorded as not parsed. Binary keystores (.p12, .pfx, .jks) are inventoried by path only; their contents are not parsed (password required)." % PEM_MAX_LINES,
    "Prose in documentation is scanned only for configuration identifiers (environment variable names, TLS directives). Narrative mentions of algorithms in release notes or user guides are not inventoried.",
    "Secret detection uses simple assignment patterns and JWT/JWK/PEM shapes. It will miss encoded or split secrets and may flag placeholder values used in tests; each Secret row is labeled with its scope (Test, CI, Documentation).",
    "The `node:lts-alpine` tag and `lts/*` CI alias resolve to different Node versions over time. The resolved version is recorded only when supplied with --lts-resolves-to.",
    "Occurrences of the same JWT literal pattern inside a single test file are collapsed to one row (first line, count in the rationale).",
    "`commit_sha` is HEAD of the checkout at scan time and `source_tree_vs_commit` records whether the scanned source matched it. Because the generated artifacts are committed on top of that source, the SHA recorded inside an artifact is always the parent of the commit that adds the artifact, never that commit itself.",
]

# --------------------------------------------------------------------------- #
# Rules
# --------------------------------------------------------------------------- #
# Each rule is a dict. Required: id, regex, category, algorithm, purpose, qclass,
# cnsa, agility, priority, rationale. Optional: key_size, mode, library,
# protocol, security_relevant, ext_dep, evidence ('match'|'label'), label,
# file_re (path filter), not_file_re, refine (callable), action, phase, gap,
# collapse_per_file (bool).

RULES: list[dict] = []


def rule(**kw):
    kw.setdefault("key_size", "")
    kw.setdefault("mode", "")
    kw.setdefault("library", "")
    kw.setdefault("protocol", "")
    kw.setdefault("security_relevant", "Yes")
    kw.setdefault("ext_dep", "")
    kw.setdefault("evidence", "match")
    kw.setdefault("action", "")
    kw.setdefault("phase", "")
    kw.setdefault("gap", "")
    # A call whose argument list may continue on the next line (`name(` then
    # whitespace, `{`, or a `[^)]` run) is matched against a short multi-line
    # window rather than a single physical line.
    # Synthetic rules have no regex; the scanner creates their rows directly
    # (for example from a parsed binary key file).
    kw.setdefault("synthetic", "regex" not in kw)
    if not kw["synthetic"]:
        kw.setdefault("multiline", bool(re.search(r"\\\((?:\\s\*|\\\{|\[\^\)\])", kw["regex"])))
        kw["regex"] = re.compile(kw["regex"])
    if "file_re" in kw:
        kw["file_re"] = re.compile(kw["file_re"])
    if "not_file_re" in kw:
        kw["not_file_re"] = re.compile(kw["not_file_re"])
    RULES.append(kw)
    return kw


def _window(lines, idx, before=WINDOW, after=WINDOW):
    lo = max(0, idx - before)
    hi = min(len(lines), idx + after + 1)
    return "\n".join(lines[lo:hi])


def _jwk_n_bits(text):
    m = re.search(r"\bn:\s*['\"]([A-Za-z0-9_\-]+)['\"]", text)
    if not m:
        return None
    n = m.group(1)
    try:
        raw = base64.urlsafe_b64decode(n + "=" * (-len(n) % 4))
    except Exception:
        return None
    return int.from_bytes(raw, "big").bit_length()


# ---- Asymmetric ---------------------------------------------------------- #

def _refine_rsa_keygen(f, lines, idx, m):
    ctx = _window(lines, idx, 2, 6)
    km = re.search(r"modulusLength:\s*(\d+)", ctx)
    if km:
        bits = int(km.group(1))
        f["key_size"] = f"RSA-{bits}"
        if bits < 2048:
            f["qclass"] = Q_WEAK
            f["rationale"] = f"RSA-{bits} is below the 2048-bit floor of SP 800-131A Rev. 2 and is quantum-vulnerable. " + f["rationale"]


rule(
    id="ASYM-RSA-KEYGEN",
    regex=r"generateKeyPair(?:Sync)?\(\s*['\"]rsa(?:-pss)?['\"]",
    category="Asymmetric", algorithm="RSA", purpose="Key pair generation for JWT signing (JWKS publication)",
    library="node:crypto", protocol="JWT/JWS", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD,
    priority=3, rationale="Generates the signing key pair used to issue tokens; algorithm and size are literals in code.",
    refine=_refine_rsa_keygen, action="Replace with configurable key type/size; add ML-DSA option when jsonwebtoken or a successor library supports it.",
    phase="Phase 3", gap="Key algorithm and modulus length are literals in code.",
)

_ASYM_KEYGEN_TYPES = {
    "ec": ("ECDSA/ECDH (EC key pair)", CNSA_SIG),
    "ed25519": ("Ed25519 (EdDSA)", CNSA_SIG),
    "ed448": ("Ed448 (EdDSA)", CNSA_SIG),
    "x25519": ("X25519 (ECDH)", CNSA_KEM),
    "x448": ("X448 (ECDH)", CNSA_KEM),
    "dsa": ("DSA", CNSA_SIG),
    "dh": ("Finite-field DH", CNSA_KEM),
}


def _refine_asym_keygen(f, lines, idx, m):
    kind = m.group(1).lower()
    f["algorithm"], f["cnsa"] = _ASYM_KEYGEN_TYPES[kind]
    # Options object of this call only: stop at the statement terminator.
    ctx = (lines[idx][m.end():] + "\n" + "\n".join(lines[idx + 1:idx + 7])).split(";", 1)[0]
    cm = re.search(r"namedCurve:\s*['\"]([\w-]+)['\"]", ctx)
    if cm:
        f["key_size"] = cm.group(1)
    elif kind in ("ed25519", "x25519"):
        f["key_size"] = "255-bit curve"
    elif kind in ("ed448", "x448"):
        f["key_size"] = "448-bit curve"
    km = re.search(r"(?:modulusLength|primeLength):\s*(\d+)", ctx)
    if km:
        f["key_size"] = f"{kind.upper()}-{km.group(1)}"


rule(
    id="ASYM-KEYGEN",
    regex=r"generateKeyPair(?:Sync)?\(\s*['\"](ec|ed25519|ed448|x25519|x448|dsa|dh)['\"]",
    category="Asymmetric", algorithm="Asymmetric key pair generation", purpose="Key pair generation",
    library="node:crypto", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=2,
    rationale="Key type is a literal at the call site.", refine=_refine_asym_keygen,
    action="Make the key type configurable; plan ML-DSA or ML-KEM replacement.", phase="Phase 3",
)


def _refine_ecdh(f, lines, idx, m):
    if m.group(1):
        f["key_size"] = m.group(1)


rule(
    id="ASYM-ECDH",
    regex=r"createECDH\(\s*(?:['\"]([\w-]+)['\"])?|\bcrypto\.diffieHellman\(|\bcomputeSecret\(",
    category="Asymmetric", algorithm="ECDH key agreement", purpose="Key exchange", library="node:crypto",
    qclass=Q_SHOR, cnsa=CNSA_KEM, agility=AGILITY_HARD, priority=2,
    rationale="Elliptic-curve key agreement is broken by Shor; the exchanged secret is exposed to harvest-now-decrypt-later.",
    refine=_refine_ecdh, action="Replace with ML-KEM (FIPS 203) or a hybrid construction.", phase="Phase 2",
)

rule(
    id="ASYM-DH",
    regex=r"createDiffieHellman(?:Group)?\(\s*(?:['\"]?(\w+)['\"]?)?|\bgetDiffieHellman\(\s*['\"](\w+)['\"]",
    category="Asymmetric", algorithm="Finite-field DH key agreement", purpose="Key exchange", library="node:crypto",
    qclass=Q_SHOR, cnsa=CNSA_KEM, agility=AGILITY_HARD, priority=2,
    rationale="Finite-field DH is broken by Shor; the exchanged secret is exposed to harvest-now-decrypt-later.",
    refine=lambda f, lines, idx, m: f.__setitem__("key_size", m.group(1) or m.group(2) or ""),
    action="Replace with ML-KEM (FIPS 203) or a hybrid construction.", phase="Phase 2",
)

rule(
    id="ASYM-SIGN-VERIFY",
    regex=r"createSign\(\s*['\"]([\w-]+)['\"]|createVerify\(\s*['\"]([\w-]+)['\"]|\bcrypto\.(sign|verify)\(",
    category="Asymmetric", algorithm="Digital signature (algorithm follows key type)", purpose="Signature generation or verification",
    library="node:crypto", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=2,
    rationale="Signature operation whose algorithm is set by the supplied key; RSA/ECDSA/EdDSA keys are quantum-vulnerable.",
    refine=lambda f, lines, idx, m: f.__setitem__("mode", (m.group(1) or m.group(2) or "").upper()),
    action="Plan ML-DSA (FIPS 204) keys once Node exposes them for this API.", phase="Phase 3",
)

rule(
    id="ASYM-WEBCRYPTO",
    regex=r"subtle\.(generateKey|importKey|sign|verify|deriveBits|deriveKey|encrypt|decrypt)\([^)]*?name:\s*['\"](ECDSA|ECDH|Ed25519|X25519|RSA-PSS|RSASSA-PKCS1-v1_5|RSA-OAEP)['\"]",
    category="Asymmetric", algorithm="Web Crypto asymmetric operation", purpose="Browser-side public-key operation",
    library="Web Crypto (crypto.subtle)", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=2,
    rationale="Public-key algorithm is a literal in browser code.",
    refine=lambda f, lines, idx, m: (f.__setitem__("algorithm", m.group(2)), f.__setitem__("mode", m.group(1)),
                                     f.__setitem__("cnsa", CNSA_KEM if m.group(2) in ("ECDH", "X25519", "RSA-OAEP") else CNSA_SIG)),
)


def _refine_jwk_private(f, lines, idx, m):
    ctx = _window(lines, idx, 2, 14)
    if re.search(r"kty:\s*['\"]RSA['\"]", ctx):
        f["algorithm"] = "RSA"
        bits = _jwk_n_bits(ctx)
        if bits:
            f["key_size"] = f"RSA-{bits}"
            if bits < 2048:
                f["qclass"] = Q_WEAK
    elif re.search(r"kty:\s*['\"]EC['\"]", ctx):
        f["algorithm"] = "ECDSA"
        cm = re.search(r"crv:\s*['\"]([\w-]+)['\"]", ctx)
        if cm:
            f["key_size"] = cm.group(1)


rule(
    id="ASYM-PRIVKEY-IMPORT",
    regex=r"createPrivateKey\(",
    category="Asymmetric", algorithm="Asymmetric private key import", purpose="Import of an embedded private key (JWK) for token signing",
    library="node:crypto", protocol="JWT/JWS", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD,
    priority=3, rationale="Private key material is embedded in source; the key reproduces a known insecure IdP default key for negative tests.",
    refine=_refine_jwk_private, evidence="label", label="crypto.createPrivateKey({... format: 'jwk'})",
    action="Keep for negative testing only; confirm the kid stays in the insecureKids denylist.", phase="Phase 1",
)


def _refine_jwk_public(f, lines, idx, m):
    ctx = _window(lines, idx, 10, 2)
    ktys = re.findall(r"kty\s*===\s*['\"](\w+)['\"]", ctx)
    if ktys:
        f["algorithm"] = "JWK import: " + "/".join(ktys) + " (RSA, ECDSA, EdDSA key types)"
        f["gap"] = "Accepted JWK key types (%s) are a literal allowlist in code; ML-DSA JWKs (kty AKP) would be dropped." % ", ".join(ktys)
    if re.search(r"use\s*===\s*['\"]sig['\"]", ctx):
        f["purpose"] += "; keys filtered to use=sig or unspecified"


rule(
    id="JWKS-KEY-IMPORT",
    regex=r"createPublicKey\(\s*\{\s*format:\s*['\"]jwk['\"]",
    category="JWT/JWS/JWKS", algorithm="JWK public key import", purpose="Import IdP JWKS keys for access-token signature verification",
    library="node:crypto", protocol="OIDC / JWKS over HTTP(S)", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD,
    ext_dep="Identity provider (OIDC) publishes JWKS; key algorithm chosen by the IdP",
    priority=1, rationale="Every API request is authenticated with a signature verified against these keys. Key type allowlist is hardcoded.",
    refine=_refine_jwk_public, action="Make the accepted key-type/algorithm allowlist configurable; add ML-DSA (kty AKP) once Node and the IdP support it.",
    phase="Phase 3",
)


def _refine_jwt_verify(f, lines, idx, m):
    ctx = _window(lines, idx, 6, 3)
    has_algs = re.search(r"\balgorithms\s*:", ctx) is not None
    has_aud = re.search(r"\baudience\s*:", ctx) is not None
    has_iss = re.search(r"\bissuer\s*:", ctx) is not None
    parts = []
    if has_algs:
        parts.append("explicit `algorithms` allowlist present")
    else:
        parts.append("no `algorithms` option: jsonwebtoken 9.x selects RS256/RS384/RS512/PS256/PS384/PS512 for RSA keys and ES256/ES384/ES512 for EC keys")
    parts.append("audience check: " + ("configurable via STIGMAN_JWT_AUD_VALUE" if has_aud else "absent"))
    parts.append("issuer check: " + ("present" if has_iss else "absent (trust is bound to the discovered JWKS)"))
    f["mode"] = "; ".join(parts)
    if not has_algs:
        f["gap"] = "Accepted JWS `alg` values are not pinned in configuration; they follow the library default for the key type."
        f["agility"] = AGILITY_LIB


rule(
    id="JWT-VERIFY",
    regex=r"\bjwt\.verify\(|jsonwebtoken\.verify\(",
    category="JWT/JWS/JWKS", algorithm="JWS signature verification (RS*/PS*/ES* per key type)", purpose="Verify bearer access tokens on API requests",
    library="jsonwebtoken", protocol="OAuth 2.0 bearer token (JWT)", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_LIB,
    ext_dep="Identity provider signing algorithm; jsonwebtoken algorithm support",
    priority=1, rationale="Primary authentication control of the API. Signature algorithm is set by the IdP and accepted by library default.",
    refine=_refine_jwt_verify, action="Pin an explicit `algorithms` allowlist from configuration; plan ML-DSA (FIPS 204) verification when the IdP and library support it.",
    phase="Phase 1 (pin) / Phase 3 (ML-DSA)",
)

rule(
    id="JWT-SIGN",
    regex=r"\bjwt\.sign\(|jsonwebtoken\.sign\(",
    category="JWT/JWS/JWKS", algorithm="JWS signing", purpose="Issue signed test tokens (mock identity provider)",
    library="jsonwebtoken", protocol="JWT/JWS", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD,
    priority=3, rationale="Token issuance exists only in the test mock; production tokens are issued by the external IdP.",
    action="Track alongside JWT-VERIFY; no production change.", phase="Phase 3",
)

rule(
    id="JWT-DECODE",
    regex=r"\bjwt\.decode\(",
    category="JWT/JWS/JWKS", algorithm="JWT parse (no verification)", purpose="Read header (kid) and payload before signature verification",
    library="jsonwebtoken", protocol="JWT", qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_NA, security_relevant="Yes",
    priority=4, rationale="Parsing only; signature is verified separately in verifyToken.",
)

CNSA_MAC = "HMAC with SHA-384 or SHA-512 and a 256-bit key (symmetric; no PQC replacement needed)"


def _set_jws_family(f, alg: str):
    """HS* is HMAC: symmetric, Grover-affected only. Everything else in the JWS
    registry that this scanner matches is an RSA/EC/EdDSA signature."""
    if alg.upper().startswith("HS") or alg == "HMAC":
        f["qclass"] = Q_GROVER
        f["cnsa"] = CNSA_MAC
        f["priority"] = 4
        f["rationale"] = "HMAC is a symmetric MAC; not affected by Shor. Key length, not algorithm, decides adequacy. " + f["rationale"]


def _refine_jwt_alg(f, lines, idx, m):
    f["algorithm"] = m.group(1)
    _set_jws_family(f, m.group(1))


rule(
    id="JWT-ALG-LITERAL",
    regex=r"\balg\s*[:=]\s*['\"](RS256|RS384|RS512|PS256|PS384|PS512|ES256|ES384|ES512|EdDSA|HS256|HS384|HS512)['\"]",
    category="JWT/JWS/JWKS", algorithm="JWS alg literal", purpose="Advertise the signing algorithm in a published JWK",
    library="jsonwebtoken / node:crypto", protocol="JWKS", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD,
    priority=3, rationale="Algorithm identifier is a literal in the test IdP.",
    refine=_refine_jwt_alg,
)

rule(
    id="JWKS-INSECURE-KID-DENYLIST",
    regex=r"insecureKids\s*=\s*\[",
    category="JWT/JWS/JWKS", algorithm="Key identifier denylist", purpose="Reject tokens signed with a known public default IdP key",
    library="application code", protocol="JWKS", qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_HARD,
    priority=4, rationale="Compensating control; the denylist is a literal in config.js, overridable only by STIGMAN_DEV_ALLOW_INSECURE_TOKENS.",
    evidence="label", label="const insecureKids = [...]",
)

rule(
    id="JWKS-ALLOW-INSECURE-FLAG",
    regex=r"allowInsecureTokens:\s*process\.env\.(\w+)",
    category="JWT/JWS/JWKS", algorithm="Configuration flag", purpose="Development override that permits known-insecure signing keys",
    library="application code", protocol="JWKS", qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_CONFIG,
    priority=2, rationale="Must remain unset in production; verify in deployment configuration.",
    action="Confirm the variable is unset in every non-development environment.", phase="Phase 1",
)

rule(
    id="JWKS-CACHE-AGE",
    regex=r"cacheMaxAge:\s*Math\.min",
    category="JWT/JWS/JWKS", algorithm="JWKS cache/rotation window", purpose="Bound how long IdP signing keys are cached (minutes; STIGMAN_JWKS_CACHE_MAX_AGE)",
    library="application code", protocol="JWKS", qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_CONFIG,
    ext_dep="Identity provider key rotation schedule",
    priority=4, rationale="Supports key rotation; unknown kids trigger a refresh (auth.js getSigningKey).",
    evidence="label", label="cacheMaxAge: Math.min(Math.max(STIGMAN_JWKS_CACHE_MAX_AGE, 1) || 10, 35791)",
)

rule(
    id="JWKS-AUDIENCE-CONFIG",
    regex=r"audienceValue:\s*process\.env\.(\w+)",
    category="JWT/JWS/JWKS", algorithm="JWT audience claim check", purpose="Optional `aud` validation for access tokens",
    library="jsonwebtoken", protocol="JWT", qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_CONFIG,
    priority=3, rationale="Audience validation is optional; recommended to set in production.",
    action="Set STIGMAN_JWT_AUD_VALUE in production.", phase="Phase 1",
)

rule(
    id="JWKS-JSON-KEY",
    regex=r"\"kty\"\s*:\s*\"(RSA|EC|OKP|oct)\"",
    category="JWT/JWS/JWKS", algorithm="JWK public key (static JWKS fixture)", purpose="Published signing key for the mock IdP used in container tests",
    library="static JSON", protocol="JWKS", qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD,
    priority=3, rationale="Static test JWKS; the key is the widely published insecure default and is on the application denylist.",
    refine=lambda f, lines, idx, m: (
        f.__setitem__("algorithm", {"RSA": "RSA", "EC": "ECDSA", "OKP": "EdDSA", "oct": "HMAC"}[m.group(1)] + " (JWK)"),
        _set_jws_family(f, "HMAC" if m.group(1) == "oct" else ""),
    ),
)

# ---- OIDC client (browser) ---------------------------------------------- #

rule(
    id="OIDC-PKCE-S256",
    regex=r"code_challenge_method['\"]?\s*,\s*['\"]S256['\"]",
    category="Hash", algorithm="SHA-256 (PKCE S256)", purpose="Bind the OAuth authorization code to the browser client",
    library="Web Crypto (crypto.subtle)", protocol="OAuth 2.0 authorization code + PKCE", qclass=Q_GROVER, cnsa=CNSA_HASH, agility=AGILITY_HARD,
    ext_dep="Identity provider must support PKCE S256",
    priority=4, rationale="SHA-256 commitment; no quantum-vulnerable public-key operation.",
)

rule(
    id="OIDC-PKCE-STRICT",
    regex=r"strictPkce\s*&&\s*!oidcConfiguration\.code_challenge_methods_supported",
    category="JWT/JWS/JWKS", algorithm="PKCE capability check", purpose="Refuse IdPs that do not advertise S256 (STIGMAN_CLIENT_STRICT_PKCE)",
    library="application code", protocol="OIDC discovery", qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_CONFIG,
    priority=4, rationale="Configuration flag; default strict.",
    evidence="label", label="ENV.strictPkce && !code_challenge_methods_supported.includes('S256')",
)

# ---- Hash ---------------------------------------------------------------- #


def _refine_hash_purpose(f, lines, idx, m):
    ctx = _window(lines, idx, 8, 8).lower()
    path = f["file"].lower()
    if "pkce" in ctx or "challenge" in ctx or "oidc-worker" in path:
        f["purpose"] = "PKCE code challenge (S256) for the OIDC authorization code flow"
        f["protocol"] = "OAuth 2.0 / PKCE"
        f["security_relevant"] = "Yes"
    elif "attachment" in ctx or "attachments" in path:
        f["purpose"] = "Digest of uploaded attachment stored as metadata (client-computed, not verified server-side)"
        f["security_relevant"] = "No (metadata / identification)"
        f["priority"] = 4
    elif "kid" in ctx or "keyid" in ctx or "key id" in ctx:
        f["purpose"] = "Derive a JWK key identifier (kid) from the DER public key"
        f["security_relevant"] = "No (identifier)"
        f["priority"] = 4
    elif "digest" in ctx and ("check" in ctx or "fix" in ctx):
        f["purpose"] = "Content digest used to detect duplicate STIG check/fix text"
        f["security_relevant"] = "No (deduplication)"
        f["priority"] = 4


rule(
    id="HASH-NODE",
    regex=r"createHash\(\s*['\"](sha1|sha256|sha384|sha512|md5|sha-1|sha-256)['\"]",
    category="Hash", algorithm="SHA-2", purpose="Digest", library="node:crypto", protocol="",
    qclass=Q_GROVER, cnsa=CNSA_HASH, agility=AGILITY_HARD, priority=4,
    rationale="Hash function literal in code.",
    refine=lambda f, lines, idx, m: (_set_hash_alg(f, m.group(1)), _refine_hash_purpose(f, lines, idx, m)),
)

rule(
    id="HASH-WEBCRYPTO",
    regex=r"subtle\.digest\(\s*['\"](SHA-1|SHA-256|SHA-384|SHA-512)['\"]",
    category="Hash", algorithm="SHA-2", purpose="Digest", library="Web Crypto (crypto.subtle)", protocol="",
    qclass=Q_GROVER, cnsa=CNSA_HASH, agility=AGILITY_HARD, priority=4,
    rationale="Hash function literal in browser code.",
    refine=lambda f, lines, idx, m: (_set_hash_alg(f, m.group(1)), _refine_hash_purpose(f, lines, idx, m)),
)


def _set_hash_alg(f, name):
    n = name.upper().replace("SHA", "SHA-").replace("SHA--", "SHA-")
    f["algorithm"] = n
    if n in ("MD5", "SHA-1"):
        f["qclass"] = Q_WEAK
        f["cnsa"] = CNSA_REMOVE
        f["priority"] = 1
        f["rationale"] = f"{n} is disallowed for digital signatures and deprecated for other integrity uses (SP 800-131A Rev. 2). " + f["rationale"]


rule(
    id="HASH-HMAC",
    regex=r"createHmac\(\s*['\"](\w+)['\"]",
    category="Symmetric", algorithm="HMAC", purpose="Message authentication", library="node:crypto",
    qclass=Q_GROVER, cnsa=CNSA_HASH, agility=AGILITY_HARD, priority=4, rationale="HMAC literal in code.",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", "HMAC-" + m.group(1).upper()),
)

rule(
    id="KDF-PASSWORD",
    regex=r"\b(bcrypt|argon2(?:id|i|d)?|pbkdf2(?:Sync)?|scrypt(?:Sync)?)\s*\(",
    category="KDF/Password", algorithm="Password hashing / KDF", purpose="Password storage or key derivation", library="",
    qclass=Q_GROVER, cnsa=CNSA_SYM, agility=AGILITY_HARD, priority=4, rationale="KDF literal in code.",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", m.group(1)),
)

# ---- Symmetric ----------------------------------------------------------- #


def _classify_symmetric(f, tok):
    tok = tok.upper()
    if "CHACHA" in tok:
        f["algorithm"] = "ChaCha20-Poly1305"
    elif "3DES" in tok or "DES-EDE" in tok or "TRIPLEDES" in tok:
        f["algorithm"] = "3DES"
        f["qclass"] = Q_WEAK
        f["cnsa"] = CNSA_SYM
        f["priority"] = 1
    else:
        f["algorithm"] = "AES"
        km = re.search(r"(128|192|256)", tok)
        if km:
            f["key_size"] = f"AES-{km.group(1)}"
            if km.group(1) != "256":
                f["rationale"] = "Below the CNSA 2.0 AES-256 requirement. " + f["rationale"]
                f["priority"] = 2
        mm = re.search(r"(GCM|CBC|CTR|ECB|CCM)", tok)
        if mm:
            f["mode"] = mm.group(1)


_CIPHERIV_CALL = re.compile(r"create(?:C|Dec)ipheriv\(")


def _refine_cipheriv(f, lines, idx, m):
    f["library"] = "node:crypto"
    literal = m.group(1)
    if not literal:
        # First argument may sit on a following line; read up to the call's
        # first closing parenthesis.
        rest = (lines[idx][m.end():] + "\n" + "\n".join(lines[idx + 1:idx + 4])).split(")", 1)[0]
        lm = re.match(r"\s*['\"]([\w-]+)['\"]", rest)
        literal = lm.group(1) if lm else None
    if literal:
        _classify_symmetric(f, literal)
    else:
        f["algorithm"] = "Symmetric cipher (algorithm supplied at runtime)"
        f["agility"] = AGILITY_CONFIG
        f["rationale"] = "Cipher name is not a literal at the call site; resolve the value in configuration."


def _refine_cipher_literal(f, lines, idx, m):
    """Skip literals that are the first argument of a createCipheriv /
    createDecipheriv call (already inventoried by SYM-CIPHERIV), whether the
    call opened on this line or on one of the three lines before it."""
    before = "\n".join(lines[max(0, idx - 3):idx] + [lines[idx][:m.start()]])
    cm = None
    for cm in _CIPHERIV_CALL.finditer(before):
        pass
    if cm is not None and ")" not in before[cm.end():]:
        return False
    _classify_symmetric(f, m.group(0))


rule(
    id="SYM-CIPHERIV",
    regex=r"create(?:C|Dec)ipheriv\(\s*(?:['\"]([\w-]+)['\"])?",
    category="Symmetric", algorithm="Symmetric cipher", purpose="Encryption / decryption", library="node:crypto",
    qclass=Q_GROVER, cnsa=CNSA_SYM, agility=AGILITY_HARD, priority=3, rationale="Cipher call in code.",
    refine=_refine_cipheriv,
)

rule(
    id="SYM-CIPHER",
    regex=r"\baes-(?:128|192|256)-(?:gcm|cbc|ctr|ecb|ccm)\b|\bAES-(?:GCM|CBC|CTR)\b|\b(?i:chacha20(?:-poly1305)?)\b|\b3DES\b|\bdes-ede3(?:-cbc)?\b|\bTripleDES\b",
    category="Symmetric", algorithm="Symmetric cipher", purpose="Encryption", library="",
    qclass=Q_GROVER, cnsa=CNSA_SYM, agility=AGILITY_HARD, priority=3, rationale="Symmetric cipher literal in code.",
    refine=_refine_cipher_literal,
)

# ---- TLS ----------------------------------------------------------------- #


def _refine_tls_server(f, lines, idx, m):
    text = "\n".join(lines)
    absent = [k for k in ("minVersion", "maxVersion", "ciphers", "secureOptions", "honorCipherOrder", "ecdhCurve") if k not in text]
    present = [k for k in ("key", "cert", "passphrase") if re.search(r"\b%s\s*[:=]" % k, text)]
    f["mode"] = "Options set: %s. Not set in code: %s (Node defaults apply)." % (", ".join(present) or "none", ", ".join(absent) or "none")
    f["gap"] = "TLS protocol versions and cipher suites are not exposed as configuration; only key/cert/passphrase file paths are."


rule(
    id="TLS-SERVER",
    regex=r"https\.createServer\(",
    category="TLS", algorithm="TLS server (native HTTPS)", purpose="Encrypt API traffic when STIGMAN_API_TLS_KEY_FILE and STIGMAN_API_TLS_CERT_FILE are set",
    library="node:https / OpenSSL", protocol="TLS (Node defaults: TLSv1.2 min, TLSv1.3 max on current Node)", qclass=Q_PROTO, cnsa=CNSA_TLS,
    agility=AGILITY_LIB, ext_dep="Operator-supplied certificate/key; browser and reverse-proxy TLS stacks; Node/OpenSSL build in the container image",
    priority=2, rationale="Key exchange in TLS is the harvest-now-decrypt-later exposure point for tokens and data in transit.",
    refine=_refine_tls_server, action="Expose minVersion (TLSv1.2 floor, prefer TLSv1.3) and cipher policy as configuration; enable hybrid X25519MLKEM768 at TLS termination in a non-production environment (Phase 2).",
    phase="Phase 1 / Phase 2",
)

rule(
    id="TLS-SERVER-PLAINTEXT-FALLBACK",
    regex=r"http\.createServer\(app\)",
    category="TLS", algorithm="Plaintext HTTP listener", purpose="Default listener when TLS files are not configured (TLS expected at a reverse proxy)",
    library="node:http", protocol="HTTP", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="Reverse proxy / ingress TLS termination", priority=2,
    rationale="Confidentiality of bearer tokens then depends entirely on an external TLS terminator.",
    action="Document the required reverse-proxy TLS policy (TLS 1.2+/1.3, approved suites) as a deployment control.", phase="Phase 1",
)

rule(
    id="TLS-API-CONFIG",
    regex=r"(key_file|cert_file|key_passphrase):\s*process\.env\.(STIGMAN_API_TLS_\w+)",
    category="TLS", algorithm="TLS server key/certificate reference", purpose="Path/passphrase configuration for native TLS",
    library="application config", protocol="TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="Operator-supplied PEM material (algorithm chosen at deployment)", priority=3,
    rationale="Certificate algorithm and key size are decided by the deployer; not visible in the repository.",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", m.group(2)),
)

rule(
    id="TLS-DB-CONFIG",
    regex=r"(ca_file|cert_file|key_file):\s*process\.env\.(STIGMAN_DB_TLS_\w+)",
    category="TLS", algorithm="MySQL TLS CA/client cert/key reference", purpose="Enable TLS (and optional client-certificate auth) to MySQL",
    library="mysql2 / node:tls", protocol="MySQL over TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="MySQL server TLS configuration and OpenSSL build", priority=3,
    rationale="Optional; database TLS is off unless one of these variables is set.",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", m.group(2)),
)


def _refine_db_ssl(f, lines, idx, m):
    text = "\n".join(lines)
    absent = [k for k in ("rejectUnauthorized", "minVersion", "ciphers") if k not in text]
    f["mode"] = "ssl object built from ca/cert/key files. Not set in code: %s (mysql2/Node defaults apply; Node verifies the server certificate by default)." % ", ".join(absent)
    f["gap"] = "MySQL TLS version and cipher policy are not configurable from the application; verification behavior relies on library defaults."


rule(
    id="TLS-DB-CLIENT",
    regex=r"poolConfig\.ssl\s*=\s*sslConfig",
    category="TLS", algorithm="TLS client (MySQL)", purpose="Encrypt database connections",
    library="mysql2 / node:tls / OpenSSL", protocol="MySQL over TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_LIB,
    ext_dep="MySQL server TLS versions and cipher suites", priority=2,
    rationale="STIG/POA&M data and credentials transit this link; key exchange is the quantum-exposed step.",
    refine=_refine_db_ssl, action="Set an explicit TLS 1.2+ floor and approved suites on the MySQL server (require_secure_transport=ON); evaluate hybrid key exchange when the server's OpenSSL supports it.",
    phase="Phase 1 / Phase 2",
)

rule(
    id="TLS-OIDC-CLIENT-CA",
    regex=r"new Agent\(\{\s*connect:\s*\{\s*ca:",
    category="TLS", algorithm="TLS client (OIDC discovery) with custom CA", purpose="Validate the IdP certificate chain using STIGMAN_OIDC_CA_CERTS",
    library="undici / node:tls", protocol="HTTPS to identity provider", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="Identity provider TLS certificate and CA", priority=3,
    rationale="Trust anchor is configurable; protocol versions follow Node defaults.",
)

rule(
    id="TLS-JWKS-CLIENT-CA",
    regex=r"requestOptions\.ca\s*=\s*this\.caCerts",
    category="TLS", algorithm="TLS client (JWKS fetch) with custom CA", purpose="Validate the IdP certificate chain when fetching JWKS",
    library="node:https", protocol="HTTPS to identity provider", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="Identity provider TLS certificate and CA", priority=3,
    rationale="Trust anchor is configurable; protocol versions follow Node defaults.",
)

rule(
    id="TLS-JWKS-HTTP-FALLBACK",
    regex=r"url\.protocol\s*===\s*['\"]https:['\"]\s*\?\s*https\s*:\s*http",
    category="TLS", algorithm="Plaintext HTTP permitted for JWKS retrieval", purpose="Fetch signing keys over http:// when the discovered jwks_uri is not https",
    library="node:http", protocol="HTTP", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="Identity provider URL scheme (STIGMAN_OIDC_PROVIDER)", priority=1,
    rationale="Signing keys fetched without TLS can be substituted on path. Default authority is http://localhost:8080 (development).",
    action="Require https for the OIDC authority outside development; verify deployment configuration.", phase="Phase 1",
    gap="Scheme is inherited from configuration with a plaintext development default.",
)

rule(
    id="OIDC-AUTHORITY-DEFAULT",
    regex=r"authority:\s*process\.env\.STIGMAN_OIDC_PROVIDER\s*\|\|\s*process\.env\.STIGMAN_API_AUTHORITY\s*\|\|\s*\"http://",
    category="TLS", algorithm="Plaintext default for OIDC authority", purpose="Development default for the identity provider URL",
    library="application config", protocol="HTTP", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="Identity provider", priority=2, rationale="Default is http://; production must override with an https URL.",
    evidence="label", label="authority: STIGMAN_OIDC_PROVIDER || STIGMAN_API_AUTHORITY || \"http://localhost:8080/realms/stigman\"",
    action="Override in every deployed environment; consider rejecting http:// authorities when not in development.", phase="Phase 1",
)

rule(
    id="TLS-DB-SERVER-DOC",
    regex=r"require_secure_transport\s*=\s*ON",
    category="TLS", algorithm="MySQL server TLS enforcement (documentation)", purpose="Documented server setting that forces TLS for all client connections",
    library="MySQL server", protocol="MySQL over TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="MySQL server", priority=3, rationale="Guidance only; enforcement happens on the database server.",
    security_relevant="Yes",
)

rule(
    id="TLS-NODE-EXTRA-CA-DOC",
    regex=r"NODE_EXTRA_CA_CERTS",
    category="TLS", algorithm="Node CA trust store extension (documentation)", purpose="Documented way to trust private CAs for outbound TLS",
    library="node:tls", protocol="TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    priority=4, rationale="Guidance only.", security_relevant="Yes", file_re=r"\.(rst|md|csv)$",
)

rule(
    id="TLS-ENV-DOC",
    regex=r"\"?(STIGMAN_API_TLS_(?:KEY_FILE|CERT_FILE|KEY_PASSPHRASE)|STIGMAN_DB_TLS_(?:CA_FILE|CERT_FILE|KEY_FILE)|STIGMAN_OIDC_CA_CERTS)\"?,\"",
    category="TLS", algorithm="TLS environment variable (documentation)", purpose="Documented TLS configuration variable",
    library="application config", protocol="TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    priority=4, rationale="Documentation of a configurable TLS input.", file_re=r"envvars\.csv$",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", m.group(1)),
)

rule(
    id="TLS-ENV-LAUNCHER",
    regex=r"^# (STIGMAN_API_TLS_(?:KEY_FILE|CERT_FILE|KEY_PASSPHRASE)|STIGMAN_DB_TLS_(?:CA_FILE|CERT_FILE|KEY_FILE)|STIGMAN_OIDC_CA_CERTS)\s*$",
    category="TLS", algorithm="TLS environment variable (launcher template)", purpose="Commented launcher template entry for a TLS variable",
    library="application config", protocol="TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    priority=4, rationale="Launcher template documents the configurable TLS inputs.", file_re=r"launchers/stig-manager\.sh$",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", m.group(1)),
)

rule(
    id="TLS-HSTS",
    regex=r"Strict-Transport-Security|\bhsts\b|helmet\(",
    category="TLS", algorithm="HSTS header", purpose="Force browsers to use HTTPS", library="", protocol="HTTP headers",
    qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG, priority=3, rationale="HSTS setting found.",
)

rule(
    id="TLS-OPTIONS-LITERAL",
    regex=r"\b(minVersion|maxVersion|ciphers|rejectUnauthorized|secureOptions|honorCipherOrder|ecdhCurve)\s*:",
    category="TLS", algorithm="Explicit TLS option", purpose="TLS version/cipher/verification option set in code", library="node:tls",
    protocol="TLS", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_HARD, priority=2, rationale="Explicit TLS option literal.",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", m.group(1)),
)

# ---- Signing of releases / trust metadata ------------------------------- #

rule(
    id="SIGN-GPG-RELEASE",
    regex=r"gpg\s+(?:--\S+\s+\S+\s+)*--armor\s+--detach-sig",
    category="Signing", algorithm="OpenPGP detached signature (key algorithm not visible in repository)", purpose="Sign release binaries in CI",
    library="GnuPG (CI runner)", protocol="OpenPGP", qclass=Q_SHOR, cnsa=CNSA_SW_SIG, agility=AGILITY_CONFIG,
    ext_dep="CI secret STIGMAN_PRIVATE_KEY; consumers' GnuPG versions", priority=2,
    rationale="Software signing is a CNSA 2.0 early-transition category; the signing key algorithm is held outside the repository.",
    evidence="label", label="gpg --armor --detach-sig <artifact>",
    action="Record the key algorithm/size of the release signing key; plan a move to LMS/XMSS or ML-DSA signatures when consumer tooling supports them.", phase="Phase 3",
)

rule(
    id="SIGN-GPG-IMPORT",
    regex=r"gpg\s+--import",
    category="Signing", algorithm="OpenPGP private key import", purpose="Load the release signing key from a CI secret",
    library="GnuPG (CI runner)", protocol="OpenPGP", qclass=Q_SHOR, cnsa=CNSA_SW_SIG, agility=AGILITY_CONFIG,
    ext_dep="CI secret store", priority=3, rationale="Key material lives in the CI secret store, not the repository.",
    evidence="label", label="gpg --import (from CI secret)",
)

rule(
    id="SIGN-GPG-VERIFY",
    regex=r"gpg\s+--verify",
    category="Signing", algorithm="OpenPGP signature verification", purpose="Verify release signatures in CI",
    library="GnuPG (CI runner)", protocol="OpenPGP", qclass=Q_SHOR, cnsa=CNSA_SW_SIG, agility=AGILITY_CONFIG,
    priority=3, rationale="Verification step for the release signing flow.", evidence="label", label="gpg --verify <sig> <artifact>",
)


def _refine_tuf_key(f, lines, idx, m):
    f["algorithm"] = "ECDSA" + (" (X.509 certificate)" if "x509" in m.group(1) else "")
    # Search from the match position so that several keys on one (minified) line
    # each pick up their own "public" value.
    ctx = m.string[m.end():] + "\n" + "\n".join(lines[idx + 1:idx + 5])
    pm = re.search(r"\"public\":\s*\"([A-Za-z0-9+/=]+)\"", ctx)
    if pm:
        info = parse_public_key_b64(pm.group(1), "x509" in m.group(1))
        if info.get("curve"):
            f["key_size"] = info["curve"]
        if info.get("subject"):
            f["mode"] = "Certificate subject: %s; valid %s to %s" % (info["subject"], info.get("not_before"), info.get("not_after"))


rule(
    id="TRUST-TUF-KEY",
    regex=r"\"keytype\":\s*\"(ecdsa(?:-x509)?)\"",
    category="Signing", algorithm="ECDSA", purpose="Container image trust metadata (Docker Content Trust / Notary root)",
    library="Notary/TUF (registry)", protocol="TUF root.json", qclass=Q_SHOR, cnsa=CNSA_SW_SIG, agility=AGILITY_HARD,
    ext_dep="Container registry Notary service and consumer docker clients", priority=2,
    rationale="Image signature trust roots are ECDSA; verification keys are pinned in this file.",
    refine=_refine_tuf_key, action="Track registry/Notary support for PQC signatures; re-issue trust roots when available.", phase="Phase 3",
    gap="Trust root keys are static ECDSA public keys embedded in root.json.",
)

# ---- Randomness ---------------------------------------------------------- #


def _refine_random(f, lines, idx, m):
    ctx = _window(lines, idx, 4, 4).lower()
    path = f["file"].lower()
    if "verifier" in ctx or "pkce" in ctx or ("oidc-worker" in path and "getrandomvalues" in m.group(0).lower()):
        f["purpose"] = "PKCE code verifier generation"
        f["security_relevant"] = "Yes"
    elif "state" in ctx or "nonce" in ctx:
        f["purpose"] = "OAuth state/nonce generation"
        f["security_relevant"] = "Yes"
    elif "channel" in ctx:
        f["purpose"] = "BroadcastChannel name"
        f["security_relevant"] = "No"
    else:
        f["purpose"] = "Identifier generation (UUID)"
        f["security_relevant"] = "No"


rule(
    id="RNG-CSPRNG",
    regex=r"crypto\.randomUUID\(|crypto\.getRandomValues\(|randomBytes\(",
    category="Randomness", algorithm="CSPRNG (platform)", purpose="Random value generation", library="node:crypto / Web Crypto",
    qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_NA, priority=4, rationale="Randomness source; not affected by Shor or Grover in a way that changes the migration plan.",
    refine=_refine_random, collapse_per_file=True,
)

# ---- Runtime ------------------------------------------------------------- #

def _refine_docker_base(f, lines, idx, m):
    value = (m.group(1) or m.group(2) or "").strip()
    if "${" in value:
        return False
    f["key_size"] = value
    if not value.startswith("node"):
        f["algorithm"] = "Container base image (build tooling; no deployed runtime crypto role)"
        f["qclass"] = Q_NA
        f["cnsa"] = CNSA_NONE
        f["priority"] = 4
        f["rationale"] = "Documentation build image; not part of the deployed system."
        f["action"] = f["gap"] = f["phase"] = ""
        f["ext_dep"] = ""
    return True


rule(
    id="RUNTIME-DOCKER-BASE",
    regex=r"^ARG\s+BASE_IMAGE\s*=\s*\"?([\w:./-]+)\"?|^FROM\s+([\w:./${}-]+)",
    category="Runtime", algorithm="Node.js/OpenSSL runtime (container base image)", purpose="Provides the TLS stack and node:crypto for the API",
    library="Container base image", protocol="", qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG,
    ext_dep="Upstream Node.js image publisher (mutable tag)", priority=2,
    rationale="Tag is mutable; the OpenSSL version that decides ML-KEM/ML-DSA availability is not pinned.",
    refine=_refine_docker_base,
    file_re=r"(^|/)Dockerfile$", action="Pin a specific Node.js LTS major (24.x) with OpenSSL 3.5+ and record it in the inventory on each release.", phase="Phase 0 / Phase 1",
    gap="Runtime version is a floating tag; PQC capability of the deployed image cannot be asserted from the repository.",
)

rule(
    id="RUNTIME-CI-NODE",
    regex=r"node-version:\s*([\w./*-]+)",
    category="Runtime", algorithm="Node.js runtime (CI)", purpose="Node used for CI test runs", library="actions/setup-node",
    qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_CONFIG, priority=4,
    rationale="CI alias resolves to the current LTS at run time.",
    refine=lambda f, lines, idx, m: f.__setitem__("key_size", m.group(1)), file_re=r"\.github/workflows/",
)

rule(
    id="RUNTIME-CI-CONTAINER-ARG",
    regex=r"build_arg:\s*\"?(node:[\w.-]+)\"?",
    category="Runtime", algorithm="Node.js runtime (container test matrix)", purpose="Base image used to build the API container in CI", library="Docker",
    qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG, priority=3, rationale="Same mutable tag as the Dockerfile default.",
    refine=lambda f, lines, idx, m: f.__setitem__("key_size", m.group(1)), file_re=r"\.github/workflows/",
)

rule(
    id="RUNTIME-PKG-TARGET",
    regex=r"\"targets\":\s*\[([^\]]+)\]",
    category="Runtime", algorithm="Node.js runtime (packaged binaries)", purpose="Node major embedded in the standalone API binaries", library="@yao-pkg/pkg",
    qclass=Q_PROTO, cnsa=CNSA_TLS, agility=AGILITY_CONFIG, priority=3,
    rationale="Binary releases embed a Node 24 runtime; the exact OpenSSL build is fixed at packaging time.",
    refine=lambda f, lines, idx, m: f.__setitem__("key_size", re.sub(r"[\"\s]", "", m.group(1))), file_re=r"pkg\.config\.json$",
)

# ---- Key material handling ---------------------------------------------- #

rule(
    id="KEYMAT-GITIGNORE",
    regex=r"^\*\*\.pem$|^api/source/tls/.*$",
    category="Certificate/Key material", algorithm="Ignored key-material paths", purpose="Keep PEM files and generated MySQL TLS config out of version control",
    library="git", protocol="", qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_NA, priority=4,
    rationale="Explains why no certificate or key files are tracked in the repository.", file_re=r"(^|/)\.gitignore$", security_relevant="Yes",
    collapse_per_file=True,
)


def _refine_pem_block(f, lines, idx, m):
    """Parse the PEM block that starts at this match and classify the row by
    the parsed strength. Public parameters only are recorded; the parsed
    metadata is also handed to the scanner (``_pem_info``) for the certificate
    table when the block is embedded in a non-certificate file."""
    kind = re.match(r"-----BEGIN ([A-Z ]+)-----", m.group(0)).group(1)
    text = lines[idx][m.start():] + "\n" + "\n".join(lines[idx + 1:idx + PEM_MAX_LINES])
    bm = _PEM_BLOCK_RE.match(text)
    info = parse_pem_block(kind, bm.group(2)) if bm else {"parser": "none", "error": f"no matching END marker within {PEM_MAX_LINES} lines"}
    f["_pem_info"] = {"source": f"PEM {kind} (embedded)", **info}
    _apply_key_info(f, info)


def _apply_key_info(f, info):
    """Copy parsed key/certificate metadata into an inventory row and classify
    the row by the parsed strength."""
    if info.get("error"):
        err = info["error"]
        f["mode"] = err if "not parsed" in err else f"not parsed: {err}"
        f["rationale"] += " Key strength not parsed; classified by file/block type only."
        return
    alg, size = info.get("key_algorithm", ""), info.get("key_size", "")
    if alg:
        f["algorithm"] += f" ({alg})"
    f["key_size"] = size
    sig = info.get("signature_algorithm", "")
    if sig:
        f["mode"] = f"signature {sig}; valid {info.get('not_before')} to {info.get('not_after')}"
    bits = re.fullmatch(r"(?:RSA|DSA)-(\d+)", size)
    if bits and int(bits.group(1)) < 2048:
        f["qclass"] = Q_WEAK
        f["rationale"] = f"{size} is below the SP 800-131A Rev. 2 floor. " + f["rationale"]
    elif re.search(r"(?i)sha1|md5|md2", sig):
        f["qclass"] = Q_WEAK
        f["rationale"] = f"Certificate signature {sig} uses a deprecated hash. " + f["rationale"]


rule(
    id="KEYMAT-PEM-PRIVATE",
    regex=r"-----BEGIN (?:RSA |EC |ENCRYPTED |OPENSSH )?PRIVATE KEY-----",
    category="Secret", algorithm="PEM private key", purpose="Embedded private key", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=1, rationale="Private key material committed to the repository.",
    evidence="label", label="[redacted] PEM private key block", refine=_refine_pem_block,
)

rule(
    id="KEYMAT-PEM-PUBLIC",
    regex=r"-----BEGIN (?:RSA )?PUBLIC KEY-----",
    category="Certificate/Key material", algorithm="PEM public key", purpose="Embedded public key", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=3, rationale="Public key committed to the repository.",
    evidence="label", label="PEM public key block", refine=_refine_pem_block,
)

rule(
    id="KEYMAT-PEM-CERT",
    regex=r"-----BEGIN CERTIFICATE-----",
    category="Certificate/Key material", algorithm="X.509 certificate (PEM)", purpose="Embedded certificate", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=3, rationale="Certificate committed to the repository.",
    evidence="label", label="PEM certificate block", refine=_refine_pem_block,
)

# Rows for binary (DER) certificate/key files and keystores are created by
# Scanner.inspect_cert_file from the parsed file, not by a regex.
rule(
    id="KEYMAT-DER-PRIVATE",
    category="Secret", algorithm="DER private key", purpose="Committed private key file", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=1, rationale="Private key material committed to the repository.",
    evidence="label", label="[redacted] DER private key file",
)

rule(
    id="KEYMAT-DER-PUBLIC",
    category="Certificate/Key material", algorithm="DER public key", purpose="Committed public key file", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=3, rationale="Public key committed to the repository.",
    evidence="label", label="DER public key file",
)

rule(
    id="KEYMAT-DER-CERT",
    category="Certificate/Key material", algorithm="X.509 certificate (DER)", purpose="Committed certificate file", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=3, rationale="Certificate committed to the repository.",
    evidence="label", label="DER certificate file",
)

rule(
    id="KEYMAT-FILE-UNPARSED",
    category="Certificate/Key material", algorithm="Key/certificate file (contents not parsed)", purpose="Committed key or certificate file", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=2,
    rationale="File with a certificate/key extension whose contents could not be parsed; it may hold private keys and must be inspected by hand.",
    evidence="label", label="key/certificate file, contents not parsed",
)

# ---- Secrets (type + location only) ------------------------------------- #

rule(
    id="SECRET-DB-PASSWORD-LITERAL",
    # Value: a quoted scalar, or an unquoted scalar up to whitespace or a
    # comment. Values starting with `$` are references, not literals.
    regex=r"\b(STIGMAN_DB_PASSWORD|MYSQL_ROOT_PASSWORD|MYSQL_PASSWORD)\b\s*[:=]\s*(?:'[^'\n]{3,}'|\"[^\"\n]{3,}\"|[^\s'\"$#][^\s#]{2,})",
    category="Secret", algorithm="Database password literal", purpose="Credential literal (type: database password)", library="",
    qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_HARD, priority=3,
    rationale="Literal credential in test/CI configuration; must not be reused in production.", evidence="label", label="[redacted] database password literal",
    security_relevant="Yes", not_file_re=r"process\.env",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", "Database password literal (%s)" % m.group(1)),
)

rule(
    id="SECRET-GENERIC-LITERAL",
    regex=r"(?i)\b(password|passwd|client_secret|api[_-]?key|secret_key|passphrase)\b['\"]?\s*[:=]\s*['\"][^'\"$]{4,}['\"]",
    category="Secret", algorithm="Credential literal", purpose="Credential literal (type from identifier name)", library="",
    qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_HARD, priority=3, rationale="Literal credential assignment.",
    evidence="label", label="[redacted] credential literal",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", "Credential literal (%s)" % m.group(1).lower()),
)

rule(
    id="SECRET-JWT-LITERAL",
    regex=r"eyJ[A-Za-z0-9_\-]{10,}\.eyJ[A-Za-z0-9_\-]{10,}\.[A-Za-z0-9_\-]{10,}",
    category="Secret", algorithm="JWT literal", purpose="Embedded signed token (type: bearer token fixture)", library="",
    qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_HARD, priority=4,
    rationale="Test fixture tokens; they are signed with test keys and expire.", evidence="label", label="[redacted] JWT literal",
    collapse_per_file=True,
)

rule(
    id="SECRET-JWK-PRIVATE",
    regex=r"\bd:\s*['\"][A-Za-z0-9_\-]{40,}['\"]",
    category="Secret", algorithm="Private key (JWK private exponent)", purpose="Embedded RSA private key component (type: JWK private key)", library="",
    qclass=Q_SHOR, cnsa=CNSA_SIG, agility=AGILITY_HARD, priority=3,
    rationale="Reproduces a public, known-insecure IdP key for negative tests; the kid is on the application denylist.",
    evidence="label", label="[redacted] JWK private component",
)

rule(
    id="SECRET-CI-REFERENCE",
    regex=r"\$\{\{\s*secrets\.([A-Z0-9_]+)\s*\}\}",
    category="Secret reference", algorithm="CI secret reference", purpose="Secret injected by the CI secret store (not stored in repository)", library="GitHub Actions",
    qclass=Q_NA, cnsa=CNSA_NONE, agility=AGILITY_CONFIG, priority=4, rationale="Reference only; value held by the CI platform.",
    refine=lambda f, lines, idx, m: f.__setitem__("algorithm", "CI secret reference (%s)" % m.group(1)), collapse_per_file=False,
)

# --------------------------------------------------------------------------- #
# Helpers: certificate / key parsing
# --------------------------------------------------------------------------- #

try:  # optional
    from cryptography import x509  # type: ignore
    from cryptography.hazmat.primitives import serialization  # type: ignore
    from cryptography.hazmat.primitives.asymmetric import ec, rsa, ed25519, ed448, dsa  # type: ignore
    HAVE_CRYPTOGRAPHY = True
except Exception:  # pragma: no cover
    HAVE_CRYPTOGRAPHY = False

HAVE_OPENSSL_CLI = shutil.which("openssl") is not None
# Distinguished-name values (CN, O, ...) may identify people or organizations.
# By default only the attribute types and the self-signed status are recorded.
INCLUDE_DN_VALUES = False


def _dn_summary(name_obj=None, text=None, other=None) -> str:
    if INCLUDE_DN_VALUES:
        return name_obj.rfc4514_string() if name_obj is not None else (text or "")
    if name_obj is not None:
        attrs = ", ".join(a.oid._name for a in name_obj)
        same = other is not None and name_obj == other
    else:
        attrs = ", ".join(re.findall(r"(?:^|[,/]\s*)([A-Za-z]+)\s*=", text or ""))
        same = other is not None and text == other
    return f"[values withheld] attributes: {attrs or 'none'}" + ("; self-signed (subject = issuer)" if same else "")


def _describe_public_key(pub) -> dict:
    if isinstance(pub, rsa.RSAPublicKey):
        return {"key_algorithm": "RSA", "key_size": f"RSA-{pub.key_size}", "curve": f"RSA-{pub.key_size}"}
    if isinstance(pub, ec.EllipticCurvePublicKey):
        return {"key_algorithm": "ECDSA", "key_size": pub.curve.name, "curve": pub.curve.name}
    if isinstance(pub, ed25519.Ed25519PublicKey):
        return {"key_algorithm": "Ed25519", "key_size": "Ed25519", "curve": "Ed25519"}
    if isinstance(pub, ed448.Ed448PublicKey):
        return {"key_algorithm": "Ed448", "key_size": "Ed448", "curve": "Ed448"}
    if isinstance(pub, dsa.DSAPublicKey):
        return {"key_algorithm": "DSA", "key_size": f"DSA-{pub.key_size}", "curve": f"DSA-{pub.key_size}"}
    return {"key_algorithm": type(pub).__name__, "key_size": "", "curve": ""}


def _parse_openssl_key_text(out: str) -> dict:
    info = {"parser": "openssl"}
    am = re.search(r"^(RSA|EC|DSA|ED25519|ED448|X25519|X448)?\s*(?:Public-Key|Private-Key):\s*(?:\((\d+) bit\))?", out, re.M)
    if am:
        alg = (am.group(1) or "").upper()
        bits = am.group(2)
        cm = re.search(r"(?:NIST CURVE|ASN1 OID): (\S+)", out)
        if cm:
            info.update({"key_algorithm": "ECDSA", "key_size": cm.group(1), "curve": cm.group(1)})
        elif alg in ("ED25519", "ED448", "X25519", "X448"):
            name = alg.capitalize()
            info.update({"key_algorithm": name, "key_size": name, "curve": name})
        elif bits:
            info.update({"key_algorithm": alg or "RSA", "key_size": f"{alg or 'RSA'}-{bits}", "curve": f"{alg or 'RSA'}-{bits}"})
    return info


def _parse_key(data: bytes, private: bool, der: bool) -> dict:
    """Algorithm and size of a public or unencrypted private key in PEM or DER
    form. Only public parameters are recorded; private components are never
    read into the output."""
    note = {"note": "private key material present; only public parameters recorded"} if private else {}
    if HAVE_CRYPTOGRAPHY:
        try:
            if der:
                key = serialization.load_der_private_key(data, password=None) if private else serialization.load_der_public_key(data)
            else:
                key = serialization.load_pem_private_key(data, password=None) if private else serialization.load_pem_public_key(data)
            info = _describe_public_key(key.public_key() if private else key)
            info["parser"] = "cryptography"
            info.update(note)
            return info
        except TypeError:
            return {"parser": "none", "error": "encrypted private key; not parsed (passphrase required)", **note}
        except Exception as e:  # noqa: BLE001
            err = str(e)
    else:
        err = "cryptography not installed"
    if HAVE_OPENSSL_CLI:
        try:
            args = ["openssl", "pkey", "-noout", "-inform", "DER" if der else "PEM"] + (["-text_pub"] if private else ["-pubin", "-text"])
            out = subprocess.run(args, input=data, capture_output=True, check=True).stdout.decode(errors="replace")
            info = _parse_openssl_key_text(out)
            if "key_algorithm" in info:
                info.update(note)
                return info
            err = "openssl output not recognised"
        except subprocess.CalledProcessError as e:
            err = f"openssl pkey could not decode the key (exit {e.returncode})"
        except Exception as e:  # noqa: BLE001
            err = str(e)
    return {"parser": "none", "error": f"key not parsed: {err}", **note}


def parse_pem_key(pem: bytes, private: bool) -> dict:
    return _parse_key(pem, private, der=False)


def parse_der_material(der: bytes) -> dict:
    """Metadata for a file with a certificate/key extension and no PEM block.
    Tries X.509 certificate, then unencrypted private key, then public key
    (PKCS#8/SPKI/PKCS#1 DER). The result carries ``material`` naming what was
    recognised; nothing recognised gives ``parser: none`` and an error."""
    info = parse_certificate_der(der)
    if not info.get("error"):
        return {"material": "certificate", **info}
    for material, private in (("private key", True), ("public key", False)):
        info = _parse_key(der, private, der=True)
        if not info.get("error") or info["error"].startswith("encrypted private key"):
            return {"material": material, **info}
    return {"parser": "none", "error": "not recognised as an X.509 certificate, private key or public key (PEM or DER)"}


def parse_certificate_der(der: bytes) -> dict:
    """Return subject/issuer/algorithms/validity for a DER certificate."""
    if HAVE_CRYPTOGRAPHY:
        try:
            cert = x509.load_der_x509_certificate(der)
            info = _describe_public_key(cert.public_key())
            sig = cert.signature_algorithm_oid._name
            info.update({
                "subject": _dn_summary(cert.subject, other=cert.issuer),
                "issuer": _dn_summary(cert.issuer, other=cert.subject),
                "signature_algorithm": sig,
                "not_before": cert.not_valid_before_utc.date().isoformat(),
                "not_after": cert.not_valid_after_utc.date().isoformat(),
                "parser": "cryptography",
            })
            return info
        except Exception as e:  # noqa: BLE001
            err = str(e)
    else:
        err = "cryptography not installed"
    if HAVE_OPENSSL_CLI:
        try:
            out = subprocess.run(["openssl", "x509", "-inform", "DER", "-noout", "-subject", "-issuer", "-dates", "-text"],
                                 input=der, capture_output=True, check=True).stdout.decode(errors="replace")
            info = {"parser": "openssl"}
            for key, pat in (("subject", r"subject=(.*)"), ("issuer", r"issuer=(.*)"), ("not_before", r"notBefore=(.*)"),
                             ("not_after", r"notAfter=(.*)"), ("signature_algorithm", r"Signature Algorithm: (\S+)"),
                             ("key_algorithm", r"Public Key Algorithm: (\S+)"), ("curve", r"(?:NIST CURVE|ASN1 OID): (\S+)|Public-Key: \((\d+ bit)\)")):
                mm = re.search(pat, out)
                if mm:
                    info[key] = next(g for g in mm.groups() if g) if mm.groups() else mm.group(0)
            info["key_size"] = info.get("curve", "")
            subj, iss = info.get("subject"), info.get("issuer")
            info["subject"] = _dn_summary(text=subj, other=iss)
            info["issuer"] = _dn_summary(text=iss, other=subj)
            return info
        except subprocess.CalledProcessError as e:
            err = f"openssl x509 could not decode the certificate (exit {e.returncode})"
        except Exception as e:  # noqa: BLE001
            err = str(e)
    return {"parser": "none", "error": err}


_PEM_BLOCK_RE = re.compile(r"-----BEGIN ([A-Z ]+)-----(.*?)-----END \1-----", re.S)


def parse_pem_block(kind: str, body: str) -> dict:
    """Metadata for one PEM block. Certificates yield subject/issuer/algorithms/
    validity; public and unencrypted private keys yield algorithm and size only.
    Encrypted keys and unknown block types are recorded as not parsed. The body
    may come from source text (quoted, `\\n`-escaped, concatenated), so only the
    base64 alphabet is kept."""
    if kind == "ENCRYPTED PRIVATE KEY" or "Proc-Type: 4,ENCRYPTED" in body:
        return {"parser": "none", "error": "encrypted private key; not parsed (passphrase required)", "key_algorithm": kind}
    b64 = re.sub(r"[^A-Za-z0-9+/=]", "", re.sub(r"\\[nrt]", "", body))
    if kind == "CERTIFICATE":
        try:
            return parse_certificate_der(base64.b64decode(b64, validate=True))
        except Exception as ex:  # noqa: BLE001
            return {"parser": "none", "error": f"malformed PEM certificate block: {ex}"}
    if "PRIVATE KEY" in kind or "PUBLIC KEY" in kind:
        wrapped = "\n".join(b64[i:i + 64] for i in range(0, len(b64), 64))
        pem = f"-----BEGIN {kind}-----\n{wrapped}\n-----END {kind}-----\n".encode()
        return parse_pem_key(pem, private="PRIVATE" in kind)
    return {"parser": "none", "error": f"{kind} block not parsed"}


def parse_public_key_b64(b64: str, is_x509: bool) -> dict:
    try:
        raw = base64.b64decode(b64)
    except Exception as e:  # noqa: BLE001
        return {"error": f"base64 decode failed: {e}"}
    if is_x509:
        pem = raw.decode(errors="replace")
        m = re.search(r"-----BEGIN CERTIFICATE-----(.*?)-----END CERTIFICATE-----", pem, re.S)
        if m:
            der = base64.b64decode(re.sub(r"\s+", "", m.group(1)))
            return parse_certificate_der(der)
        return parse_certificate_der(raw)
    if HAVE_CRYPTOGRAPHY:
        try:
            return _describe_public_key(serialization.load_der_public_key(raw))
        except Exception as e:  # noqa: BLE001
            return {"error": str(e)}
    if HAVE_OPENSSL_CLI:
        try:
            out = subprocess.run(["openssl", "pkey", "-pubin", "-inform", "DER", "-text", "-noout"], input=raw,
                                 capture_output=True, check=True).stdout.decode(errors="replace")
            mm = re.search(r"(?:NIST CURVE|ASN1 OID): (\S+)|Public-Key: \((\d+) bit\)", out)
            curve = (mm.group(1) or f"RSA-{mm.group(2)}") if mm else ""
            return {"curve": curve, "key_size": curve, "parser": "openssl"}
        except Exception as e:  # noqa: BLE001
            return {"error": str(e)}
    return {"error": "no parser available"}


# --------------------------------------------------------------------------- #
# Scanner
# --------------------------------------------------------------------------- #

def scope_for(rel: str) -> str:
    if rel.startswith("test/") or "/test/" in rel or rel.endswith(".test.js"):
        return "Test"
    if rel.startswith(".github/"):
        return "CI"
    if rel.startswith("docs/") or rel.endswith((".rst", ".md")):
        return "Documentation"
    if rel.startswith("client/"):
        return "Client (browser)"
    if rel.startswith("api/"):
        return "API (server)"
    return "Repository"


def iter_files(repo: Path, extra_exclude_dirs: set[str] = frozenset(), extra_exclude_files: set[str] = frozenset()):
    exclude_paths = EXCLUDE_DIR_PATHS | set(extra_exclude_dirs)
    exclude_files = EXCLUDE_FILE_PATHS | set(extra_exclude_files)
    for root, dirs, files in os.walk(repo):
        rel_root = Path(root).relative_to(repo).as_posix()
        dirs[:] = sorted(d for d in dirs if d not in EXCLUDE_DIR_NAMES
                         and (d if rel_root == "." else f"{rel_root}/{d}") not in exclude_paths)
        for name in sorted(files):
            p = Path(root) / name
            rel = p.relative_to(repo).as_posix()
            if rel in exclude_files:
                continue
            yield p, rel


def is_text_candidate(p: Path, rel: str) -> bool:
    if any(r.search(rel) for r in SKIP_FILE_RES):
        return False
    if p.name in TEXT_BASENAMES or p.name.startswith("Dockerfile"):
        return True
    if p.suffix.lower() in TEXT_EXTENSIONS:
        return True
    return p.suffix == "" and p.stat().st_size < 512 * 1024


def read_text(p: Path):
    try:
        if p.stat().st_size > MAX_TEXT_BYTES:
            return None
        data = p.read_bytes()
    except OSError:
        return None
    if b"\x00" in data[:4096]:
        return None
    return data.decode("utf-8", errors="replace")


class Scanner:
    def __init__(self, repo: Path, lts_resolves_to: str | None, probe_node: bool, out_dir: Path | None = None,
                 out_files: tuple[Path, ...] = ()):
        self.repo = repo
        self.lts_resolves_to = lts_resolves_to
        self.probe_node = probe_node
        # Generated artifacts inside the repository are never scanned: the
        # output directory as a whole when it is a subdirectory, and the
        # concrete output files always (the directory cannot be excluded when
        # it is the repository root itself).
        self.extra_exclude_dirs: set[str] = set()
        self.extra_exclude_files: set[str] = set()
        if out_dir is not None and out_dir.resolve().is_relative_to(repo):
            rel_out = out_dir.resolve().relative_to(repo).as_posix()
            if rel_out != ".":
                self.extra_exclude_dirs.add(rel_out)
        for f in out_files:
            if f.resolve().is_relative_to(repo):
                self.extra_exclude_files.add(f.resolve().relative_to(repo).as_posix())
        self.findings: list[dict] = []
        self.certificates: list[dict] = []
        self.libraries: list[dict] = []
        self.runtime: dict = {}
        self.files_scanned = 0
        self.files_skipped = 0
        self.lockfiles: list[str] = []
        self.manifests: list[str] = []

    # -- main ---------------------------------------------------------------
    def run(self):
        for p, rel in iter_files(self.repo, self.extra_exclude_dirs, self.extra_exclude_files):
            if p.name in LOCKFILE_NAMES:
                self.lockfiles.append(rel)
                continue
            if p.name == "package.json":
                self.manifests.append(rel)
            if p.suffix.lower() in CERT_EXTENSIONS:
                pem_text = self.inspect_cert_file(p, rel)
                if pem_text is not None:
                    self.files_scanned += 1
                    self.scan_lines(rel, pem_text.split("\n"))
                continue
            if not is_text_candidate(p, rel):
                self.files_skipped += 1
                continue
            text = read_text(p)
            if text is None:
                self.files_skipped += 1
                continue
            self.files_scanned += 1
            lines = text.split("\n")
            self.scan_lines(rel, lines)
            if '"kty"' in text and '"keys"' in text:
                self.inspect_jwks_json(rel, text)
        self.collect_libraries()
        self.collect_runtime()
        self.finalize()

    # -- pattern scan -------------------------------------------------------
    def scan_lines(self, rel: str, lines: list[str]):
        scope = scope_for(rel)
        seen_collapse: dict[str, dict] = {}
        for r in RULES:
            if r["synthetic"] or ("file_re" in r and not r["file_re"].search(rel)):
                continue
            for idx, line in enumerate(lines):
                if "not_file_re" in r and r["not_file_re"].search(line):
                    continue
                # Call rules see this line plus the next few, so an argument on
                # the following line still matches; only matches that start on
                # this line are kept so each call is inventoried once.
                target = "\n".join(lines[idx:idx + CALL_WINDOW]) if r["multiline"] else line
                # finditer: minified JSON/YAML can hold several assets on one line.
                for m in r["regex"].finditer(target):
                    if m.start() > len(line):
                        break
                    f = self.new_finding(r, rel, idx + 1, m, scope)
                    if r.get("refine") and r["refine"](f, lines, idx, m) is False:
                        continue
                    pem_info = f.pop("_pem_info", None)
                    if pem_info is not None and Path(rel).suffix.lower() not in CERT_EXTENSIONS:
                        self.certificates.append({"file": rel, "line": idx + 1, "scope": scope, **pem_info})
                    if r.get("collapse_per_file"):
                        key = f"{r['id']}|{f['purpose']}"
                        if key in seen_collapse:
                            seen_collapse[key]["_count"] += 1
                            continue
                        seen_collapse[key] = f
                    self.findings.append(f)
        for f in seen_collapse.values():
            if f["_count"] > 1:
                f["rationale"] += f" {f['_count']} occurrences in this file; first occurrence listed."

    def new_finding(self, r: dict, rel: str, line: int, m, scope: str) -> dict:
        evidence = r["label"] if r["evidence"] == "label" else " ".join(m.group(0).split())[:90]
        f = {
            "asset_id": "",
            "file": rel,
            "line": line,
            "scope": scope,
            "category": r["category"],
            "algorithm": r["algorithm"],
            "key_size": r["key_size"],
            "mode": r["mode"],
            "purpose": r["purpose"],
            "library": r["library"],
            "protocol": r["protocol"],
            "security_relevant": r["security_relevant"],
            "qclass": r["qclass"],
            "cnsa": r["cnsa"],
            "agility": r["agility"],
            "ext_dep": r["ext_dep"],
            "priority": r["priority"],
            "rationale": r["rationale"],
            "action": r["action"],
            "phase": r["phase"],
            "gap": r["gap"],
            "evidence": evidence,
            "rule": r["id"],
            "_count": 1,
        }
        return f

    # -- JWKS JSON ----------------------------------------------------------
    def inspect_jwks_json(self, rel: str, text: str):
        try:
            data = json.loads(text)
        except Exception:
            return
        keys = data.get("keys") if isinstance(data, dict) else None
        if not isinstance(keys, list):
            return
        # JWKS-JSON-KEY rows are appended in document order (finditer per line),
        # and json.loads preserves array order, so the i-th structured key that
        # the rule can match pairs with the i-th row for this file. This holds
        # for pretty-printed and single-line JWKS alike.
        rows = [f for f in self.findings if f["rule"] == "JWKS-JSON-KEY" and f["file"] == rel]
        structured = [k for k in keys if isinstance(k, dict) and k.get("kty") in ("RSA", "EC", "OKP", "oct")]
        paired = list(zip(structured, rows)) if len(structured) == len(rows) else [(k, None) for k in structured]
        for k, f in paired:
            kid = k.get("kid", "")
            line = f["line"] if f else 1
            bits = None
            if k.get("kty") == "RSA" and k.get("n"):
                n = k["n"]
                try:
                    bits = int.from_bytes(base64.urlsafe_b64decode(n + "=" * (-len(n) % 4)), "big").bit_length()
                except Exception:
                    bits = None
            if f is not None:
                if bits:
                    f["key_size"] = f"RSA-{bits}"
                    if bits < 2048:
                        f["qclass"] = Q_WEAK
                        f["rationale"] = f"RSA-{bits} is below the SP 800-131A Rev. 2 floor. " + f["rationale"]
                elif k.get("kty") in ("EC", "OKP") and k.get("crv"):
                    f["key_size"] = k["crv"]
                if k.get("alg"):
                    f["algorithm"] += f", alg={k['alg']}"
            for j, c in enumerate(k.get("x5c") or []):
                try:
                    info = parse_certificate_der(base64.b64decode(c))
                except Exception as e:  # noqa: BLE001
                    info = {"parser": "none", "error": str(e)}
                info.update({"file": rel, "line": line, "source": f"JWKS x5c[{j}] (kid {kid})", "scope": scope_for(rel)})
                self.certificates.append(info)

    # -- certificate files --------------------------------------------------
    def inspect_cert_file(self, p: Path, rel: str) -> str | None:
        """Record certificate/key metadata. Returns the PEM text when the file is
        text so the caller can also run the pattern rules (KEYMAT-*) over it;
        returns None for binary files."""
        entry = {"file": rel, "line": 1, "source": f"file ({p.suffix})", "scope": scope_for(rel)}
        try:
            data = p.read_bytes()
        except OSError as ex:
            entry.update({"parser": "none", "error": f"unreadable: {type(ex).__name__}"})
            self.certificates.append(entry)
            self.add_file_finding("KEYMAT-FILE-UNPARSED", rel, entry)
            self.files_skipped += 1
            return None
        if p.suffix.lower() in (".p12", ".pfx", ".jks"):
            entry.update({"parser": "none", "error": "binary keystore; not parsed (password required)"})
            self.certificates.append(entry)
            self.add_file_finding("KEYMAT-FILE-UNPARSED", rel, entry)
            return None
        text = data.decode("utf-8", errors="replace")
        blocks = [(m.group(1), m.group(2), text.count("\n", 0, m.start()) + 1) for m in _PEM_BLOCK_RE.finditer(text)]
        if not blocks:
            try:
                info = parse_der_material(data)
            except Exception as e:  # noqa: BLE001
                info = {"parser": "none", "error": f"not parsed: {e}"}
            material = info.pop("material", None)
            if material:
                entry["source"] += f" DER {material}"
            entry.update(info)
            self.certificates.append(entry)
            rule_id = {"certificate": "KEYMAT-DER-CERT", "private key": "KEYMAT-DER-PRIVATE",
                       "public key": "KEYMAT-DER-PUBLIC"}.get(material, "KEYMAT-FILE-UNPARSED")
            self.add_file_finding(rule_id, rel, entry)
            return None
        for kind, body, line in blocks:
            e = dict(entry)
            e["line"] = line
            e["source"] = f"PEM {kind}"
            e.update(parse_pem_block(kind, body))
            self.certificates.append(e)
        return text

    def add_file_finding(self, rule_id: str, rel: str, info: dict):
        """Inventory row for a whole binary key/certificate file, classified by
        the parsed metadata like an embedded PEM block."""
        r = next(r for r in RULES if r["id"] == rule_id)
        f = self.new_finding(r, rel, 1, None, scope_for(rel))
        _apply_key_info(f, info)
        self.findings.append(f)

    # -- libraries ----------------------------------------------------------
    def collect_libraries(self):
        declared: dict[str, dict] = {}
        for rel in self.manifests:
            try:
                pkg = json.loads((self.repo / rel).read_text())
            except Exception:
                continue
            for section in ("dependencies", "devDependencies", "optionalDependencies"):
                for name, spec in (pkg.get(section) or {}).items():
                    if name in CRYPTO_LIBRARIES:
                        declared.setdefault(name, {})[rel] = f"{spec} ({section})"
        resolved: dict[str, dict] = {}
        for rel in self.lockfiles:
            try:
                lock = json.loads((self.repo / rel).read_text())
            except Exception:
                continue
            packages = lock.get("packages") or {}
            for path, meta in packages.items():
                name = path.split("node_modules/")[-1] if "node_modules/" in path else None
                if name in CRYPTO_LIBRARIES and isinstance(meta, dict) and meta.get("version"):
                    resolved.setdefault(name, {}).setdefault(rel, set()).add(meta["version"])
        for name, meta in CRYPTO_LIBRARIES.items():
            if name not in declared and name not in resolved:
                continue
            self.libraries.append({
                "library": name,
                "declared": "; ".join(f"{k}: {v}" for k, v in sorted(declared.get(name, {}).items())) or "transitive only",
                "resolved": "; ".join(f"{k}: {', '.join(sorted(v))}" for k, v in sorted(resolved.get(name, {}).items())) or "not in lockfile",
                "role": meta["role"],
                "pqc_status": meta["pqc"],
                "notes": meta["notes"],
            })
        # Built-in and platform crypto
        node_crypto_files = sorted({f["file"] for f in self.findings if f["library"].startswith("node:") or "node:crypto" in f["library"]})
        webcrypto_files = sorted({f["file"] for f in self.findings if "Web Crypto" in f["library"]})
        self.libraries.append({
            "library": "node:crypto / node:tls / node:https (built-in)",
            "declared": "Node.js runtime",
            "resolved": self.runtime_summary_placeholder(),
            "role": "JWK import, RSA key generation (tests), SHA-256 digests, TLS server and clients",
            "pqc_status": "ML-KEM and ML-DSA APIs are available in Node.js 24.x builds with OpenSSL 3.5+ (see Runtime probe). jsonwebtoken does not use them.",
            "notes": "Used in: " + ", ".join(node_crypto_files),
        })
        self.libraries.append({
            "library": "Web Crypto API (browser crypto.subtle / getRandomValues)",
            "declared": "Browser platform",
            "resolved": "Set by the user's browser",
            "role": "PKCE S256 digest, random verifier/state/nonce, attachment digest",
            "pqc_status": "No public-key operations performed in the client; TLS to API and IdP is provided by the browser.",
            "notes": "Used in: " + ", ".join(webcrypto_files),
        })

    def runtime_summary_placeholder(self):
        return "see Runtime"

    # -- runtime ------------------------------------------------------------
    def collect_runtime(self):
        rt = {"pins": [], "local_probe": None, "lts_resolves_to": self.lts_resolves_to}
        for f in self.findings:
            if f["category"] == "Runtime" and f["key_size"].startswith("node"):
                rt["pins"].append({"file": f["file"], "line": f["line"], "source": f["algorithm"], "value": f["key_size"]})
        if self.probe_node and shutil.which("node"):
            script = r"""
const crypto = require('node:crypto'); const tls = require('node:tls');
const out = { node: process.versions.node, openssl: process.versions.openssl,
  tls_default_min: tls.DEFAULT_MIN_VERSION, tls_default_max: tls.DEFAULT_MAX_VERSION };
for (const alg of ['ml-kem-768','ml-dsa-65']) {
  try { crypto.generateKeyPairSync(alg); out[alg] = 'available'; } catch (e) { out[alg] = 'unavailable: ' + e.message; }
}
console.log(JSON.stringify(out));
"""
            try:
                res = subprocess.run(["node", "-e", script], capture_output=True, text=True, timeout=30)
                rt["local_probe"] = json.loads(res.stdout.strip().splitlines()[-1]) if res.returncode == 0 else {"error": res.stderr.strip()[:300]}
            except Exception as e:  # noqa: BLE001
                rt["local_probe"] = {"error": str(e)}
        self.runtime = rt
        for lib in self.libraries:
            if lib["resolved"] == "see Runtime":
                probe = rt["local_probe"] or {}
                lib["resolved"] = ("Repository pins: " + "; ".join(f"{p['value']} ({p['file']})" for p in rt["pins"] if p["value"]))
                if probe.get("node"):
                    lib["resolved"] += f". Scanning host: Node {probe['node']} / OpenSSL {probe['openssl']}"
                if self.lts_resolves_to:
                    lib["resolved"] += f". lts tag resolution supplied: {self.lts_resolves_to}"

    # -- absence findings ---------------------------------------------------
    def add_absence_findings(self):
        anchor = next((f for f in self.findings if f["rule"] == "TLS-SERVER"), None)
        if not anchor:
            return
        matched = {f["rule"] for f in self.findings}
        absences = [
            ("TLS-HSTS", "ABSENT-HSTS", "HSTS header (not configured)",
             "No Strict-Transport-Security header or helmet() use in application code; browsers are not instructed to require HTTPS by the API itself.",
             "Reverse proxy / ingress must set HSTS", 3, "Set HSTS at the TLS terminator; document it as a deployment control."),
            ("TLS-OPTIONS-LITERAL", "ABSENT-TLS-OPTIONS", "Explicit TLS options (none set)",
             "No minVersion, maxVersion, ciphers, rejectUnauthorized, secureOptions, honorCipherOrder or ecdhCurve literal anywhere in scanned files. Node.js defaults apply to the API server, the MySQL client and the IdP clients.",
             "Node.js/OpenSSL build defaults in the container image", 2, "Add configuration for a TLS 1.2 floor / TLS 1.3 preference and an approved cipher policy; verify hybrid ML-KEM group availability in the deployed image."),
        ]
        for src_rule, rule_id, algorithm, mode, ext_dep, prio, action in absences:
            if src_rule in matched:
                continue
            self.findings.append({
                "asset_id": "", "file": anchor["file"], "line": anchor["line"], "scope": anchor["scope"], "category": "TLS",
                "algorithm": algorithm, "key_size": "", "mode": mode, "purpose": "Absence check across all scanned files",
                "library": "node:tls", "protocol": "TLS", "security_relevant": "Yes", "qclass": Q_PROTO, "cnsa": CNSA_TLS,
                "agility": AGILITY_LIB, "ext_dep": ext_dep, "priority": prio,
                "rationale": "Absence recorded so the control can be assigned to the deployment layer or added as configuration.",
                "action": action, "phase": "Phase 1", "gap": mode, "evidence": "pattern not found in any scanned file", "rule": rule_id, "_count": 1,
            })

    # -- finalize -----------------------------------------------------------
    def finalize(self):
        self.add_absence_findings()
        for f in self.findings:
            if f["scope"] == "Test" and f["category"] not in ("Secret",):
                if f["priority"] < 3:
                    f["priority"] = 3
                f["rationale"] += " Test scope: not deployed."
            if f["scope"] == "Documentation" and f["priority"] < 3:
                f["priority"] = 3
            f.pop("_count", None)
        cat_order = {c: i for i, c in enumerate([
            "JWT/JWS/JWKS", "Asymmetric", "TLS", "Signing", "Hash", "Symmetric", "KDF/Password",
            "Certificate/Key material", "Runtime", "Randomness", "Secret", "Secret reference"])}
        self.findings.sort(key=lambda f: (cat_order.get(f["category"], 99), f["file"], f["line"], f["rule"]))
        for i, f in enumerate(self.findings, 1):
            f["asset_id"] = f"CI-{i:04d}"

    # -- derived tables -----------------------------------------------------
    def dependencies(self) -> list[dict]:
        groups: OrderedDict[tuple, dict] = OrderedDict()
        for f in self.findings:
            if not f["ext_dep"]:
                continue
            key = (f["ext_dep"], f["algorithm"].split(" (")[0])
            g = groups.setdefault(key, {"component": f"{f['file']} ({f['scope']})", "external_party": f["ext_dep"],
                                        "algorithm": f["algorithm"], "protocol": f["protocol"], "qclass": f["qclass"],
                                        "cnsa": f["cnsa"], "assets": []})
            g["assets"].append(f["asset_id"])
        rows = []
        for g in groups.values():
            g["assets"] = ", ".join(g["assets"])
            rows.append(g)
        return rows

    def agility_gaps(self) -> list[dict]:
        rows = []
        for f in self.findings:
            if f["agility"] in (AGILITY_LIB, AGILITY_HARD) and f["qclass"] in (Q_SHOR, Q_WEAK, Q_PROTO) and f["scope"] != "Test":
                rows.append({"asset_id": f["asset_id"], "file": f["file"], "line": f["line"], "algorithm": f["algorithm"],
                             "agility": f["agility"], "gap": f["gap"] or "Algorithm or key type is fixed in code or by the library default.",
                             "recommendation": f["action"] or "Expose as configuration or track library support.", "phase": f["phase"] or "Phase 3"})
        return rows

    def migration_candidates(self) -> list[dict]:
        rows = []
        for f in sorted(self.findings, key=lambda x: (x["priority"], x["asset_id"])):
            if f["priority"] <= 2 or f["qclass"] == Q_WEAK:
                rows.append({"asset_id": f["asset_id"], "priority": f["priority"], "file": f["file"], "line": f["line"],
                             "scope": f["scope"], "algorithm": f["algorithm"], "qclass": f["qclass"], "cnsa": f["cnsa"],
                             "ext_dep": f["ext_dep"], "action": f["action"] or "Track; no change until dependency supports PQC.",
                             "phase": f["phase"] or "Phase 3", "rationale": f["rationale"]})
        return rows

    def summary(self, commit: str, generated: str, tree_state: str) -> dict:
        by_class = Counter(f["qclass"] for f in self.findings)
        by_priority = Counter(f["priority"] for f in self.findings)
        by_category = Counter(f["category"] for f in self.findings)
        by_scope = Counter(f["scope"] for f in self.findings)
        by_agility = Counter(f["agility"] for f in self.findings)
        return {
            "generated_utc": generated,
            "commit_sha": commit,
            "source_tree_vs_commit": tree_state,
            "files_scanned": self.files_scanned,
            "files_skipped": self.files_skipped,
            "rules": len(RULES),
            "total_findings": len(self.findings),
            "by_quantum_class": OrderedDict((c, by_class.get(c, 0)) for c in CLASS_ORDER),
            "by_priority": OrderedDict((p, by_priority.get(p, 0)) for p in (1, 2, 3, 4)),
            "by_category": OrderedDict(sorted(by_category.items())),
            "by_scope": OrderedDict(sorted(by_scope.items())),
            "by_agility": OrderedDict(sorted(by_agility.items())),
            "certificates_parsed": len([c for c in self.certificates if not c.get("error") and c.get("parser") not in (None, "none")]),
            "certificates_unparsed": len([c for c in self.certificates if c.get("error") or c.get("parser") in (None, "none")]),
            "libraries": len(self.libraries),
        }


# --------------------------------------------------------------------------- #
# Output: JSON
# --------------------------------------------------------------------------- #

INVENTORY_COLUMNS = [
    ("asset_id", "Asset ID"), ("file", "File"), ("line", "Line"), ("scope", "Scope"), ("category", "Category"),
    ("algorithm", "Algorithm"), ("key_size", "Key size / curve"), ("mode", "Mode / options"), ("purpose", "Purpose"),
    ("library", "Library"), ("protocol", "Protocol context"), ("security_relevant", "Security-relevant"),
    ("qclass", "Quantum-vulnerability class"), ("cnsa", "CNSA 2.0 target"), ("agility", "Crypto-agility"),
    ("ext_dep", "External dependency"), ("priority", "Migration priority"), ("rationale", "Rationale"),
    ("action", "Recommended action"), ("phase", "Roadmap phase"), ("evidence", "Evidence (matched text)"), ("rule", "Rule"),
]


def git_commit(repo: Path) -> str:
    try:
        return subprocess.run(["git", "-C", str(repo), "rev-parse", "HEAD"], capture_output=True, text=True, check=True).stdout.strip()
    except Exception:
        return "unknown"


def git_tree_state(repo: Path, ignore_dirs: set[str], ignore_files: set[str] = frozenset()) -> str:
    """Whether the scanned source matches HEAD. The output artifacts are ignored
    because they are always newer than the commit they record: they are
    committed on top of the SHA they were generated from."""
    try:
        out = subprocess.run(["git", "-C", str(repo), "status", "--porcelain", "--untracked-files=all"],
                             capture_output=True, text=True, check=True).stdout
    except Exception:
        return "unknown (git status unavailable)"
    ignored = {"docs/security/crypto-inventory"} | set(ignore_dirs)
    changed = [ln[3:] for ln in out.splitlines() if ln.strip()
               and ln[3:] not in ignore_files
               and not any(ln[3:].startswith(d + "/") for d in ignored)
               and not EXCLUDE_DIR_NAMES.intersection(ln[3:].split("/")[:-1])]
    return "clean: scanned source equals HEAD" if not changed else f"modified: {len(changed)} source path(s) differ from HEAD"


def build_document(scanner: Scanner, repo: Path) -> dict:
    generated = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    commit = git_commit(repo)
    tree_state = git_tree_state(repo, scanner.extra_exclude_dirs, scanner.extra_exclude_files)
    return OrderedDict([
        ("schema_version", "1.0"),
        ("summary", scanner.summary(commit, generated, tree_state)),
        ("method", {"checks": METHOD_CHECKS, "limits": METHOD_LIMITS,
                    "excluded_dirs": sorted(EXCLUDE_DIR_NAMES | EXCLUDE_DIR_PATHS | scanner.extra_exclude_dirs),
                    "excluded_files": sorted(EXCLUDE_FILE_PATHS | scanner.extra_exclude_files),
                    "parsers": {"cryptography": HAVE_CRYPTOGRAPHY, "openssl_cli": HAVE_OPENSSL_CLI}}),
        ("runtime", scanner.runtime),
        ("libraries", scanner.libraries),
        ("certificates", scanner.certificates),
        ("dependencies", scanner.dependencies()),
        ("crypto_agility_gaps", scanner.agility_gaps()),
        ("migration_candidates", scanner.migration_candidates()),
        ("inventory", scanner.findings),
    ])


# --------------------------------------------------------------------------- #
# Output: XLSX
# --------------------------------------------------------------------------- #

def write_xlsx(doc: dict, path: Path):
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        print("openpyxl not installed; skipping XLSX output", file=sys.stderr)
        return False

    wb = Workbook()
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(title, headers, rows, widths=None, table_name=None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = header_font
            c.alignment = wrap
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(48, len(h) + 6)))
        ws.freeze_panes = "A2"
        if table_name and rows:
            ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
            t = Table(displayName=table_name, ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        elif rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        return ws

    # Inventory
    inv = doc["inventory"]
    headers = [h for _, h in INVENTORY_COLUMNS]
    rows = [[f[k] for k, _ in INVENTORY_COLUMNS] for f in inv]
    widths = {"File": 42, "Algorithm": 34, "Mode / options": 48, "Purpose": 46, "Rationale": 60, "Recommended action": 50,
              "Quantum-vulnerability class": 34, "CNSA 2.0 target": 40, "External dependency": 40, "Evidence (matched text)": 40,
              "Key size / curve": 18, "Library": 22, "Protocol context": 28}
    ws = sheet("Inventory", headers, rows, widths, table_name="Inventory")
    if rows:
        col = get_column_letter(headers.index("Quantum-vulnerability class") + 1)
        last = get_column_letter(len(headers))
        rng = f"A2:{last}{len(rows) + 1}"
        fills = {
            Q_SHOR: "F8CBAD", Q_WEAK: "FF9999", Q_GROVER: "FFE699", Q_PROTO: "BDD7EE", Q_NA: "E2EFDA",
        }
        for cls, color in fills.items():
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col}2="{cls}"'], fill=PatternFill("solid", fgColor=color), stopIfTrue=False))

    # Libraries
    sheet("Libraries", ["Library", "Declared (manifest: range)", "Resolved version (lockfile)", "Crypto role", "PQC support status today", "Notes"],
          [[lib["library"], lib["declared"], lib["resolved"], lib["role"], lib["pqc_status"], lib["notes"]] for lib in doc["libraries"]],
          {"Library": 30, "Declared (manifest: range)": 34, "Resolved version (lockfile)": 44, "Crypto role": 44, "PQC support status today": 60, "Notes": 60},
          table_name="Libraries")

    # TLS and certificates
    tls_rows = [[f["asset_id"], f["file"], f["line"], f["scope"], f["algorithm"], f["mode"], f["purpose"], f["ext_dep"], f["qclass"], f["priority"]]
                for f in inv if f["category"] == "TLS"]
    cert_rows = [[c.get("file"), c.get("line"), c.get("scope"), c.get("source"), c.get("subject", ""), c.get("issuer", ""),
                  c.get("key_algorithm", ""), c.get("key_size", ""), c.get("signature_algorithm", ""), c.get("not_before", ""),
                  c.get("not_after", ""), c.get("parser", ""), c.get("error", "")] for c in doc["certificates"]]
    ws = sheet("TLS-and-Certificates", ["Asset ID", "File", "Line", "Scope", "TLS item", "Options / observations", "Purpose", "External dependency", "Quantum-vulnerability class", "Priority"],
               tls_rows, {"File": 40, "TLS item": 40, "Options / observations": 60, "Purpose": 50, "External dependency": 40, "Quantum-vulnerability class": 34}, table_name="TLSConfig")
    start = len(tls_rows) + 4
    ws.cell(row=start - 1, column=1, value="Certificates and public keys found in the tree").font = header_font
    cert_headers = ["File", "Line", "Scope", "Source", "Subject", "Issuer", "Key algorithm", "Key size / curve", "Signature algorithm", "Not before", "Not after", "Parser", "Error / note"]
    for i, h in enumerate(cert_headers, 1):
        c = ws.cell(row=start, column=i, value=h)
        c.font = header_font
    if cert_rows:
        for r_i, r in enumerate(cert_rows, start + 1):
            for c_i, v in enumerate(r, 1):
                ws.cell(row=r_i, column=c_i, value=v).alignment = wrap
    else:
        ws.cell(row=start + 1, column=1, value="No certificate or key files (%s) are tracked in the repository." % ", ".join(sorted(CERT_EXTENSIONS)))

    # Dependencies
    sheet("Dependencies", ["System component (file, scope)", "External party", "Algorithm", "Protocol", "Quantum-vulnerability class", "CNSA 2.0 target", "Inventory assets"],
          [[d["component"], d["external_party"], d["algorithm"], d["protocol"], d["qclass"], d["cnsa"], d["assets"]] for d in doc["dependencies"]],
          {"System component (file, scope)": 44, "External party": 50, "Algorithm": 40, "Protocol": 30, "Quantum-vulnerability class": 34, "CNSA 2.0 target": 44, "Inventory assets": 30},
          table_name="Dependencies")

    # Gaps
    sheet("Crypto-Agility-Gaps", ["Asset ID", "File", "Line", "Algorithm", "Agility rating", "Gap", "Recommendation", "Phase"],
          [[g["asset_id"], g["file"], g["line"], g["algorithm"], g["agility"], g["gap"], g["recommendation"], g["phase"]] for g in doc["crypto_agility_gaps"]],
          {"File": 40, "Algorithm": 40, "Gap": 60, "Recommendation": 60}, table_name="AgilityGaps")

    # Migration candidates
    sheet("Migration-Candidates", ["Asset ID", "Priority", "File", "Line", "Scope", "Algorithm", "Quantum-vulnerability class", "CNSA 2.0 target", "External dependency", "Recommended action", "Phase", "Rationale"],
          [[m["asset_id"], m["priority"], m["file"], m["line"], m["scope"], m["algorithm"], m["qclass"], m["cnsa"], m["ext_dep"], m["action"], m["phase"], m["rationale"]] for m in doc["migration_candidates"]],
          {"File": 40, "Algorithm": 40, "Quantum-vulnerability class": 34, "CNSA 2.0 target": 40, "External dependency": 40, "Recommended action": 60, "Rationale": 60}, table_name="MigrationCandidates")

    # Summary
    s = doc["summary"]
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 60
    kv = [("Generated (UTC)", s["generated_utc"]), ("Commit SHA scanned", s["commit_sha"]), ("Source tree vs. commit", s["source_tree_vs_commit"]),
          ("Files pattern-scanned", s["files_scanned"]),
          ("Files skipped (binary/oversize/excluded type)", s["files_skipped"]), ("Detection rules", s["rules"]), ("Inventory rows", s["total_findings"]),
          ("Certificate/key entries parsed", s["certificates_parsed"]), ("Certificate/key entries not parsed", s["certificates_unparsed"]), ("Libraries inventoried", s["libraries"])]
    for k, v in kv:
        ws.append([k, v])
    for label, d in (("Rows by quantum-vulnerability class", s["by_quantum_class"]), ("Rows by migration priority (1 = highest)", s["by_priority"]),
                     ("Rows by category", s["by_category"]), ("Rows by scope", s["by_scope"]), ("Rows by crypto-agility rating", s["by_agility"])):
        ws.append([])
        ws.append([label, "Count"])
        ws.cell(row=ws.max_row, column=1).font = header_font
        ws.cell(row=ws.max_row, column=2).font = header_font
        for k, v in d.items():
            ws.append([str(k), v])
    ws.append([])
    ws.append(["Runtime pins found", ""])
    ws.cell(row=ws.max_row, column=1).font = header_font
    for p in doc["runtime"]["pins"]:
        ws.append([f"{p['file']}:{p['line']} ({p['source']})", p["value"]])
    probe = doc["runtime"].get("local_probe")
    if probe:
        ws.append([])
        ws.append(["Local runtime probe (scanning host, not the deployed image)", ""])
        ws.cell(row=ws.max_row, column=1).font = header_font
        for k, v in probe.items():
            ws.append([k, str(v)])
    if doc["runtime"].get("lts_resolves_to"):
        ws.append(["node:lts / lts/* resolution supplied at scan time", doc["runtime"]["lts_resolves_to"]])

    # Method
    ws = wb.create_sheet("Method")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 140
    ws.append(["Type", "Description"])
    for c in ws[1]:
        c.font = header_font
    for c in doc["method"]["checks"]:
        ws.append(["Check", c])
    for lim in doc["method"]["limits"]:
        ws.append(["Limit", lim])
    ws.append(["Excluded", "Directories not scanned: " + ", ".join(doc["method"]["excluded_dirs"])])
    ws.append(["Parsers", "cryptography available: %s; openssl CLI available: %s" % (doc["method"]["parsers"]["cryptography"], doc["method"]["parsers"]["openssl_cli"])])
    ws.append(["Classes", "Quantum-vulnerability classes: " + " | ".join(CLASS_ORDER)])
    ws.append(["Priority", "1 = act in Phase 1 (hygiene) or blocks later phases; 2 = Phase 1/2 configuration and TLS work; 3 = track, depends on external party or test-only; 4 = informational, no migration action."])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap

    del wb["Sheet"]
    wb.save(path)
    return True


# --------------------------------------------------------------------------- #
# Output: Markdown report from template
# --------------------------------------------------------------------------- #

def md_table(headers, rows):
    def esc(v):
        return str(v).replace("|", "\\|").replace("\n", " ")
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for r in rows:
        out.append("| " + " | ".join(esc(v) for v in r) + " |")
    return "\n".join(out)


Q_SHORT = {Q_SHOR: "Shor", Q_GROVER: "Grover", Q_WEAK: "Weak", Q_PROTO: "Protocol", Q_NA: "N/A"}


def render_report(doc: dict, template: Path, out: Path):
    s = doc["summary"]
    inv = doc["inventory"]
    top = sorted(inv, key=lambda f: (f["priority"], 0 if f["scope"] != "Test" else 1, f["asset_id"]))[:10]

    def loc(f):
        return f"{f['file']}:{f['line']}"

    ctx = {
        "COMMIT_SHA": s["commit_sha"],
        "SOURCE_TREE_STATE": s["source_tree_vs_commit"],
        "GENERATED_UTC": s["generated_utc"],
        "TOTAL_ROWS": s["total_findings"],
        "FILES_SCANNED": s["files_scanned"],
        "RULE_COUNT": s["rules"],
        "COUNT_SHOR": s["by_quantum_class"][Q_SHOR],
        "COUNT_WEAK": s["by_quantum_class"][Q_WEAK],
        "COUNT_GROVER": s["by_quantum_class"][Q_GROVER],
        "COUNT_PROTO": s["by_quantum_class"][Q_PROTO],
        "COUNT_NA": s["by_quantum_class"][Q_NA],
        "COUNT_P1": s["by_priority"][1], "COUNT_P2": s["by_priority"][2], "COUNT_P3": s["by_priority"][3], "COUNT_P4": s["by_priority"][4],
        "TABLE_BY_CLASS": md_table(["Quantum-vulnerability class", "Rows"], list(s["by_quantum_class"].items())),
        "TABLE_BY_PRIORITY": md_table(["Migration priority", "Rows"], list(s["by_priority"].items())),
        "TABLE_BY_CATEGORY": md_table(["Category", "Rows"], list(s["by_category"].items())),
        "TABLE_BY_SCOPE": md_table(["Scope", "Rows"], list(s["by_scope"].items())),
        "TABLE_BY_AGILITY": md_table(["Crypto-agility rating", "Rows"], list(s["by_agility"].items())),
        "TABLE_TOP10": md_table(["Asset", "Location", "Algorithm", "Class", "P", "Rationale"],
                                [[f["asset_id"], loc(f), f["algorithm"], Q_SHORT[f["qclass"]], f["priority"], f["rationale"]] for f in top]),
        "TABLE_LIBRARIES": md_table(["Library", "Resolved version", "Crypto role", "PQC support today"],
                                    [[lib["library"], lib["resolved"], lib["role"], lib["pqc_status"]] for lib in doc["libraries"]]),
        "TABLE_TLS": md_table(["Asset", "Location", "Item", "Observations"],
                              [[f["asset_id"], loc(f), f["algorithm"], f["mode"] or f["purpose"]] for f in inv if f["category"] == "TLS" and f["scope"] not in ("Documentation",)]),
        "TABLE_CERTS": md_table(["Location", "Source", "Key / signature", "Valid to", "Subject"],
                                [[f"{c['file']}:{c['line']}", re.sub(r"\(kid ([A-Za-z0-9_-]{12})[A-Za-z0-9_-]+\)", r"(kid \1…)", c.get("source", "")), f"{c.get('key_size', '')} / {c.get('signature_algorithm', '')}", c.get("not_after", ""), c.get("subject", c.get("error", ""))] for c in doc["certificates"]]) if doc["certificates"] else "No certificate material found.",
        "TABLE_DEPENDENCIES": md_table(["Assets", "Component", "External party", "Algorithm / item", "Class"],
                                       [[d["assets"], d["component"].split(" (")[0], d["external_party"], d["algorithm"], Q_SHORT[d["qclass"]]] for d in doc["dependencies"]]),
        "TABLE_GAPS": md_table(["Asset", "Location", "Algorithm", "Rating", "Gap"],
                               [[g["asset_id"], loc(g), g["algorithm"], g["agility"], g["gap"]] for g in doc["crypto_agility_gaps"]]),
        "TABLE_CANDIDATES": md_table(["Asset", "P", "Location", "Algorithm", "Action", "Phase"],
                                     [[m["asset_id"], m["priority"], loc(m), m["algorithm"], m["action"], m["phase"]] for m in doc["migration_candidates"]]),
        "TABLE_SECRETS": md_table(["Asset", "Location", "Scope", "Type"],
                                  [[f["asset_id"], loc(f), f["scope"], f["algorithm"]] for f in inv if f["category"] == "Secret"]),
        "TABLE_RUNTIME": md_table(["Location", "Source", "Value"], [[f"{p['file']}:{p['line']}", p["source"], p["value"]] for p in doc["runtime"]["pins"]]),
        "RUNTIME_PROBE": ("; ".join(f"{k}: {v}" for k, v in doc["runtime"]["local_probe"].items()) if doc["runtime"].get("local_probe") else "not probed"),
        "LTS_RESOLVES_TO": doc["runtime"].get("lts_resolves_to") or "not supplied",
        "COUNT_SHOR_NONTEST": sum(1 for f in inv if f["qclass"] == Q_SHOR and f["scope"] != "Test"),
        "COUNT_SHOR_TEST": sum(1 for f in inv if f["qclass"] == Q_SHOR and f["scope"] == "Test"),
        "COUNT_WEAK_NONTEST": sum(1 for f in inv if f["qclass"] == Q_WEAK and f["scope"] != "Test"),
        "COUNT_TEST_SCOPE": s["by_scope"].get("Test", 0),
        "COUNT_SECRETS": s["by_category"].get("Secret", 0),
        "TABLE_SHOR": md_table(["Asset", "Location", "Algorithm", "Key size/curve", "Agility", "P"],
                               [[f["asset_id"], loc(f), f["algorithm"], f["key_size"], f["agility"], f["priority"]] for f in inv if f["qclass"] == Q_SHOR]),
        "TABLE_WEAK": md_table(["Asset", "Location", "Algorithm", "Key size/curve", "Rationale"],
                               [[f["asset_id"], loc(f), f["algorithm"], f["key_size"], f["rationale"]] for f in inv if f["qclass"] == Q_WEAK]),
        "TABLE_GROVER": md_table(["Asset", "Location", "Algorithm", "Purpose", "Security-relevant"],
                                 [[f["asset_id"], loc(f), f["algorithm"], f["purpose"], f["security_relevant"]] for f in inv if f["qclass"] == Q_GROVER]),
        "TABLE_SECRETS_BY_TYPE": md_table(["Scope", "Type", "Rows"],
                                          [[k[0], k[1], v] for k, v in sorted(Counter((f["scope"], f["algorithm"]) for f in inv if f["category"] == "Secret").items())]),
        "METHOD_LIMITS": "\n".join(f"- {lim}" for lim in METHOD_LIMITS),
        "METHOD_CHECKS": "\n".join(f"- {c}" for c in METHOD_CHECKS),
    }
    text = template.read_text()
    missing = set(re.findall(r"\{\{(\w+)\}\}", text)) - set(ctx)
    if missing:
        raise SystemExit(f"template placeholders without values: {sorted(missing)}")
    for k, v in ctx.items():
        text = text.replace("{{%s}}" % k, str(v))
    out.write_text(text)


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--repo", default=str(Path(__file__).resolve().parents[1]), help="repository root (default: parent of scripts/)")
    ap.add_argument("--out", required=True, help="output directory")
    ap.add_argument("--json-name", default="crypto-inventory.json")
    ap.add_argument("--xlsx-name", default="Cryptographic-Inventory.xlsx")
    ap.add_argument("--report-template", default=None, help="markdown template with {{PLACEHOLDERS}} (default: <out>/PQC-Readiness-Assessment.template.md if present)")
    ap.add_argument("--report-name", default="PQC-Readiness-Assessment.md")
    ap.add_argument("--lts-resolves-to", default=None, help="record what node:lts-alpine / lts/* resolved to at scan time (e.g. '24.21.0, OpenSSL 3.5.x')")
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--no-probe", action="store_true", help="do not probe the local node binary")
    ap.add_argument("--dn-values", action="store_true", help="record full certificate subject/issuer values (default: attribute types only)")
    args = ap.parse_args(argv)
    global INCLUDE_DN_VALUES
    INCLUDE_DN_VALUES = args.dn_values

    repo = Path(args.repo).resolve()
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)

    template = Path(args.report_template).resolve() if args.report_template else out / "PQC-Readiness-Assessment.template.md"
    out_files = (out / args.json_name, out / args.xlsx_name, out / args.report_name, template)
    scanner = Scanner(repo, args.lts_resolves_to, probe_node=not args.no_probe, out_dir=out, out_files=out_files)
    scanner.run()
    doc = build_document(scanner, repo)

    json_path = out / args.json_name
    json_path.write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n")
    print(f"wrote {json_path} ({doc['summary']['total_findings']} inventory rows)")

    if not args.no_xlsx:
        xlsx_path = out / args.xlsx_name
        if write_xlsx(doc, xlsx_path):
            print(f"wrote {xlsx_path}")

    if template.exists():
        render_report(doc, template, out / args.report_name)
        print(f"wrote {out / args.report_name}")
    else:
        print(f"no report template at {template}; skipped report rendering")

    s = doc["summary"]
    print("rows by class: " + "; ".join(f"{k}: {v}" for k, v in s["by_quantum_class"].items()))
    print("rows by priority: " + "; ".join(f"P{k}: {v}" for k, v in s["by_priority"].items()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
