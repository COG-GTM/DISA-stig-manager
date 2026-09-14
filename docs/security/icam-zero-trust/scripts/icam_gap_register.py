#!/usr/bin/env python3
"""Generate ICAM-ZT-Gap-Register.xlsx from icam_gap_data.py.

Usage:
    python3 icam_gap_register.py            # write the workbook, print sheet row counts
    python3 icam_gap_register.py --check    # also verify ICAM-ZT-Gap-Assessment.md is in sync

Requires openpyxl. The workbook is fully derived from icam_gap_data.py; do
not hand-edit it. ``--check`` fails (exit 1) when the Markdown assessment
references a gap ID that is not in the data, omits a gap ID that is, or states
risk counts that differ from the data.
"""

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import icam_gap_data as data  # noqa: E402

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE.parent
XLSX_PATH = OUT_DIR / "ICAM-ZT-Gap-Register.xlsx"
MD_PATH = OUT_DIR / "ICAM-ZT-Gap-Assessment.md"

HEADER_FILL = PatternFill("solid", fgColor="3969CA")
HEADER_FONT = Font(bold=True, color="FAFAFA")
RISK_FILLS = {
    "H": "F4B6B6",
    "M": "FBE3A8",
    "L": "CDEFD9",
}
WRAP = Alignment(wrap_text=True, vertical="top")

GAP_KEYS = [
    "id", "affected", "pillar", "source", "current", "gap", "risk",
    "action", "authority", "dependency", "effort", "phase",
]
assert len(GAP_KEYS) == len(data.GAP_COLUMNS)


def risk_level(gap):
    return gap["risk"].split(" ", 1)[0].strip("—- ")


def validate_data():
    ids = [g["id"] for g in data.GAPS]
    dupes = [i for i, c in Counter(ids).items() if c > 1]
    if dupes:
        raise SystemExit(f"duplicate gap ids: {dupes}")
    for g in data.GAPS:
        missing = [k for k in GAP_KEYS if not g.get(k)]
        if missing:
            raise SystemExit(f"{g['id']}: empty fields {missing}")
        if g["pillar"] not in data.PILLARS:
            raise SystemExit(f"{g['id']}: unknown pillar {g['pillar']!r}")
        if risk_level(g) not in data.RISK_LEVELS:
            raise SystemExit(f"{g['id']}: risk must start with H/M/L: {g['risk'][:20]!r}")
        if g["effort"] not in data.EFFORT_LEVELS:
            raise SystemExit(f"{g['id']}: effort must be S/M/L")
        if g["phase"] not in data.PHASES:
            raise SystemExit(f"{g['id']}: unknown phase {g['phase']!r}")
        if ":" not in g["affected"] and "Needs Government input" not in g["source"]:
            raise SystemExit(f"{g['id']}: affected must cite file:line or source must be 'Needs Government input'")
    for row in data.CROSSWALK:
        if row[2] not in data.CROSSWALK_STATUSES:
            raise SystemExit(f"crosswalk {row[0]}: bad status {row[2]!r}")
        for ref in filter(None, [r.strip() for r in row[4].split(",")]):
            if ref not in ids:
                raise SystemExit(f"crosswalk {row[0]}: unknown gap {ref}")
    phases = [r[0] for r in data.ROADMAP]
    if tuple(phases) != data.PHASES:
        raise SystemExit("ROADMAP phases must match PHASES in order")
    pillars = [r[0] for r in data.PILLAR_SUMMARY]
    if pillars != data.PILLARS:
        raise SystemExit("PILLAR_SUMMARY must cover PILLARS in order")


def style_header(ws, ncols, widths):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(ws.max_row, 1)}"
    ws.row_dimensions[1].height = 45


def wrap_all(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def build_gaps(wb):
    ws = wb.active
    ws.title = "Gaps"
    ws.append(data.GAP_COLUMNS)
    for g in data.GAPS:
        ws.append([g[k] for k in GAP_KEYS])
    widths = [8, 42, 16, 40, 60, 50, 40, 60, 30, 40, 8, 12]
    style_header(ws, len(data.GAP_COLUMNS), widths)
    wrap_all(ws)
    last = ws.max_row
    risk_col = get_column_letter(GAP_KEYS.index("risk") + 1)
    rng = f"A2:{get_column_letter(len(GAP_KEYS))}{last}"
    for level, color in RISK_FILLS.items():
        ws.conditional_formatting.add(
            rng,
            FormulaRule(formula=[f'LEFT(${risk_col}2,1)="{level}"'], fill=PatternFill("solid", fgColor=color)),
        )
    return ws


def build_crosswalk(wb):
    ws = wb.create_sheet("Control-Crosswalk")
    ws.append(["NIST SP 800-53r5 control", "Title", "Status", "Evidence", "Related gaps"])
    for row in data.CROSSWALK:
        ws.append(list(row))
    style_header(ws, 5, [14, 40, 22, 80, 18])
    wrap_all(ws)
    fills = {
        "Satisfied": "CDEFD9",
        "Partially": "FBE3A8",
        "Not satisfied": "F4B6B6",
        "Needs Government input": "D9E2F3",
    }
    rng = f"A2:E{ws.max_row}"
    for status, color in fills.items():
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'$C2="{status}"'], fill=PatternFill("solid", fgColor=color))
        )
    return ws


def build_pillars(wb):
    ws = wb.create_sheet("ZT-Pillar-Summary")
    ws.append(["ZT pillar", "As-is", "Target", "Gap count", "High", "Medium", "Low", "Gap IDs"])
    for pillar, asis, target in data.PILLAR_SUMMARY:
        gaps = [g for g in data.GAPS if g["pillar"] == pillar]
        counts = Counter(risk_level(g) for g in gaps)
        ws.append([
            pillar, asis, target, len(gaps),
            counts.get("H", 0), counts.get("M", 0), counts.get("L", 0),
            ", ".join(g["id"] for g in gaps),
        ])
    style_header(ws, 8, [22, 60, 60, 10, 8, 8, 8, 40])
    wrap_all(ws)
    return ws


def build_roadmap(wb):
    ws = wb.create_sheet("Roadmap")
    ws.append(["Phase", "Name", "Activities", "Dependencies", "Government decision needed", "Gaps addressed"])
    for phase, name, acts, deps, decision in data.ROADMAP:
        gaps = ", ".join(g["id"] for g in data.GAPS if g["phase"] == phase)
        ws.append([phase, name, acts, deps, decision, gaps])
    style_header(ws, 6, [10, 28, 70, 45, 55, 40])
    wrap_all(ws)
    return ws


def build_method(wb):
    ws = wb.create_sheet("Method")
    ws.append(["Item", "Description"])
    for k, v in data.METHOD:
        ws.append([k, v])
    counts = Counter(risk_level(g) for g in data.GAPS)
    ws.append(["Gap count (derived)", f"{len(data.GAPS)} total: H={counts.get('H', 0)}, M={counts.get('M', 0)}, L={counts.get('L', 0)}"])
    style_header(ws, 2, [24, 120])
    wrap_all(ws)
    ws2 = wb.create_sheet("Sources")
    ws2.append(["Source ID", "Title", "URL (verified) or blank", "Verification note"])
    for s in data.SOURCES:
        ws2.append([s["id"], s["title"], s["url"] or "", s["verified"]])
    style_header(ws2, 4, [14, 70, 60, 70])
    wrap_all(ws2)
    return ws


def generate():
    validate_data()
    wb = Workbook()
    build_gaps(wb)
    build_crosswalk(wb)
    build_pillars(wb)
    build_roadmap(wb)
    build_method(wb)
    wb.save(XLSX_PATH)
    wb2 = load_workbook(XLSX_PATH)
    print(f"wrote {XLSX_PATH}")
    for name in wb2.sheetnames:
        ws = wb2[name]
        print(f"  {name}: {ws.max_row - 1} data rows, {ws.max_column} columns")
    counts = Counter(risk_level(g) for g in data.GAPS)
    print(f"gaps by risk: H={counts.get('H', 0)} M={counts.get('M', 0)} L={counts.get('L', 0)} total={len(data.GAPS)}")
    return counts


def check_markdown(counts):
    if not MD_PATH.exists():
        raise SystemExit(f"--check: {MD_PATH} not found")
    text = MD_PATH.read_text(encoding="utf-8")
    data_ids = {g["id"] for g in data.GAPS}
    md_ids = set(re.findall(r"\bG-\d{2}\b", text))
    problems = []
    if md_ids - data_ids:
        problems.append(f"markdown references unknown gaps: {sorted(md_ids - data_ids)}")
    if data_ids - md_ids:
        problems.append(f"markdown omits gaps: {sorted(data_ids - md_ids)}")
    expected = f"{len(data.GAPS)} gaps ({counts.get('H', 0)} High, {counts.get('M', 0)} Medium, {counts.get('L', 0)} Low)"
    if expected not in text:
        problems.append(f"markdown does not contain the derived count string: {expected!r}")
    problems += check_crosswalk_rows(text)
    problems += check_pillar_rows(text)
    problems += check_roadmap_rows(text)
    if problems:
        for p in problems:
            print("CHECK FAILED:", p)
        raise SystemExit(1)
    print(
        f"--check passed: {len(data_ids)} gap ids, count string, {len(data.CROSSWALK)} crosswalk rows, "
        f"{len(data.PILLARS)} pillar rows and {len(data.ROADMAP)} roadmap phases synchronised with {MD_PATH.name}"
    )


def gap_ref_set(text):
    return set(re.findall(r"\bG-\d{2}\b", text))


def single_match(matches, what):
    """Return (match, problem): exactly one Markdown row may describe a given control, pillar or phase."""
    if not matches:
        return None, f"markdown omits {what}"
    if len(matches) > 1:
        return None, f"markdown has {len(matches)} rows for {what}; expected exactly one"
    return matches[0], None


def check_crosswalk_rows(text):
    """Each NIST control row in the Markdown table must carry the data file's status and gap references."""
    problems = []
    md_rows = {}
    for m in re.finditer(r"^\| ((?:IA|AC|AU|SC)-[\d()]+) \| ([^|]+) \| ([^|]*) \| ([^|]*) \|$", text, re.M):
        md_rows.setdefault(m.group(1), []).append((m.group(2).strip(), gap_ref_set(m.group(4))))
    for control, _title, status, _evidence, gaps in data.CROSSWALK:
        row, problem = single_match(md_rows.get(control, []), f"crosswalk row for control {control}")
        if problem:
            problems.append(problem)
            continue
        md_status, md_gaps = row
        if md_status != status:
            problems.append(f"crosswalk {control}: markdown status {md_status!r} != data {status!r}")
        if md_gaps != gap_ref_set(gaps):
            problems.append(f"crosswalk {control}: markdown gaps {sorted(md_gaps)} != data {sorted(gap_ref_set(gaps))}")
    extra = set(md_rows) - {row[0] for row in data.CROSSWALK}
    if extra:
        problems.append(f"markdown crosswalk table has controls absent from data: {sorted(extra)}")
    return problems


def check_pillar_rows(text):
    """The gap list in each pillar row header must equal the gaps assigned to that pillar in the data."""
    problems = []
    for pillar in data.PILLARS:
        matches = re.findall(r"^\| \*\*" + re.escape(pillar) + r"\*\* \(([^)]*)\) \|", text, re.M)
        expected = {g["id"] for g in data.GAPS if g["pillar"] == pillar}
        m, problem = single_match(matches, f"pillar row for {pillar}")
        if problem:
            problems.append(problem)
            continue
        found = gap_ref_set(m)
        if found != expected:
            problems.append(f"pillar {pillar}: markdown gaps {sorted(found)} != data {sorted(expected)}")
    return problems


def check_roadmap_rows(text):
    """The 'Gaps:' list of each roadmap phase bullet must equal the gaps sequenced into that phase in the data."""
    problems = []
    for phase, *_ in data.ROADMAP:
        matches = re.findall(r"^- \*\*" + re.escape(phase) + r" — .*?Gaps: ([^\n]*)$", text, re.M)
        expected = {g["id"] for g in data.GAPS if g["phase"] == phase}
        m, problem = single_match(matches, f"roadmap bullet for {phase}")
        if problem:
            problems.append(problem)
            continue
        found = gap_ref_set(m)
        if found != expected:
            problems.append(f"roadmap {phase}: markdown gaps {sorted(found)} != data {sorted(expected)}")
    return problems


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true", help="verify the Markdown assessment matches the data")
    args = ap.parse_args()
    counts = generate()
    if args.check:
        check_markdown(counts)


if __name__ == "__main__":
    main()
